# -*- coding: utf-8 -*-
"""Tests du resolveur de dimensions (plan BA de reference, cotes en cm).

Chaines de cotation du plan :
  X : 248 + 725 + 238 + 478 + 265 + 460 = 2414
  Y : 382 +  40 +  82 + 108 +  15 + 382 = 1009
L'axe 3 (X) n'est pas cote : il se deduit de 725 - 238.
"""
from __future__ import annotations

import math

import pytest

from core.dimension_resolver import (
    AxisGraph,
    Cote,
    DimensionResolver,
    ElementSpec,
    Method,
    ResolverConfig,
    Status,
    calibrate_scale,
    cotations_from_number_centers,
    format_fr,
    overlay_label,
    parse_direct,
    polygon_area,
    polygon_area_with_holes,
    polyline_area,
)

# ---------------------------------------------------------------------------
# Donnees du plan (positions = cumul en cm, echelle 1 pour le graphe)
# ---------------------------------------------------------------------------
X_CHAIN = [(248, 0, 248), (725, 248, 973), (238, 973, 1211),
           (478, 1211, 1689), (265, 1689, 1954), (460, 1954, 2414)]
Y_CHAIN = [(382, 0, 382), (40, 382, 422), (82, 422, 504),
           (108, 504, 612), (15, 612, 627), (382, 627, 1009)]
X_AXES = ["1", "2", "4", "5", "6", "7", "8"]
Y_AXES = ["A", "B", "C", "D", "E", "F", "G"]


def _resolver(scale: bool = False) -> DimensionResolver:
    r = DimensionResolver(
        ResolverConfig(unit="cm"),
        axis_names={"X": X_AXES, "Y": Y_AXES},
    )
    r.add_cotes([Cote(value=v, start=a, end=b, direction="X", label=f"{v}")
                 for v, a, b in X_CHAIN])
    r.add_cotes([Cote(value=v, start=a, end=b, direction="Y", label=f"{v}")
                 for v, a, b in Y_CHAIN])
    if scale:
        r.calibrate()
    return r


# ---------------------------------------------------------------------------
# 1. Graphe 1D : axes cumules
# ---------------------------------------------------------------------------
def test_axes_x_cumules():
    axes = _resolver().axes("X")
    positions = [round(a.position, 3) for a in axes]
    assert positions == [0, 248, 973, 1211, 1689, 1954, 2414]
    assert [a.name for a in axes] == X_AXES
    assert all(a.determined for a in axes)


def test_axes_y_cumules():
    axes = _resolver().axes("Y")
    assert [round(a.position, 3) for a in axes] == [0, 382, 422, 504, 612, 627, 1009]


def test_chaine_x_somme_exacte():
    """"""
    r = _resolver()
    span = r.resolve_span("X", 0, 2414)
    assert span.value == pytest.approx(2414)
    assert span.method is Method.SOMME
    assert span.status is Status.OK
    assert span.formula == "=248+725+238+478+265+460"


def test_emprise_2414_x_1009():
    r = _resolver()
    res = r.resolve_element(ElementSpec(
        label="Emprise", category="emprise", element_id="EMP",
        axis_start="1", axis_end="8", direction="X",
        pos_start=0, pos_end=2414, width=1009))
    assert res.L == pytest.approx(2414)
    assert res.l == pytest.approx(1009)
    assert res.surface_m2 == pytest.approx(243.57, abs=0.01)
    assert res.method is Method.SOMME
    assert res.status is Status.OK


# ---------------------------------------------------------------------------
# 2. M3a : axe 3 non cote
# ---------------------------------------------------------------------------
def test_axe_3_sous_cote():
    """2->3 + 3->4 doit valoir exactement 725 ; 3->4 mesure = 238."""
    r = _resolver()
    r.add_cote(Cote(value=238, start=735, end=973, direction="X", label="3-4"))
    missing = r.resolve_missing("X", 248, 973)
    assert missing.value == pytest.approx(487)
    assert missing.method is Method.SOUSTRACTION
    assert missing.formula == "=725-238"
    assert missing.value + 238 == pytest.approx(725)


def test_axe_3_position_via_cote_manquante():
    r = _resolver()
    r.add_cote(Cote(value=238, start=735, end=973, direction="X", label="3-4"))
    axes = {a.name: a.position for a in r.axes("X")}
    # l'axe 3 n'a pas de nom : on verifie la position du noeud a 735 cm
    assert 735 in [round(p, 3) for p in axes.values()] or any(
        abs(p - 735) < 2 for p in axes.values())


# ---------------------------------------------------------------------------
# 3. M3b : longueurs nettes
# ---------------------------------------------------------------------------
def test_longueur_nette_6_7():
    """265 - 60 - 60 = 145 (deux semelles S3 de 120)."""
    r = _resolver()
    res = r.resolve_element(ElementSpec(
        label="CH axe 6-7", category="chainage", element_id="CH1",
        axis_start="6", axis_end="7", direction="X", pos_start=1689, pos_end=1954,
        width_start=120, width_end=120, declared_text="CH(20x30)"))
    assert res.mode == "net"
    assert res.L == pytest.approx(145)
    assert res.method is Method.SOUSTRACTION
    assert res.status is Status.OK
    assert res.formula == "=265-120/2-120/2"


def test_longueur_nette_4_5():
    """238 - 75 - 55 = 108 (S4 150 et S2 110)."""
    r = _resolver()
    res = r.resolve_element(ElementSpec(
        label="LG axe 4-5", category="longrine", element_id="LG1",
        axis_start="4", axis_end="5", direction="X", pos_start=973, pos_end=1211,
        width_start=150, width_end=110, declared_text="LG(30x40)"))
    assert res.L == pytest.approx(108)
    assert res.formula == "=238-150/2-110/2"
    assert res.volume_m3 == pytest.approx(108 * 30 * 40 * 1e-6, abs=1e-4)


def test_mode_brut_impose():
    r = DimensionResolver(ResolverConfig(unit="cm", length_mode="brut"))
    r.add_cotes([Cote(value=v, start=a, end=b, direction="X") for v, a, b in X_CHAIN])
    res = r.resolve_element(ElementSpec(
        label="LG axe 4-5", category="longrine", axis_start="4", axis_end="5",
        pos_start=973, pos_end=1211, width_start=150, width_end=110))
    assert res.mode == "brut"
    assert res.L == pytest.approx(238)


def test_net_sans_largeur_reste_brut():
    r = _resolver()
    res = r.resolve_element(ElementSpec(
        label="CH axe 6-7", category="chainage", pos_start=1689, pos_end=1954))
    assert res.mode == "brut"
    assert res.L == pytest.approx(265)


# ---------------------------------------------------------------------------
# 4. M1 : lecture directe
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("text,expected", [
    ("S4(150x150x40)", (150, 150, 40)),
    ("150 x 150 x 40", (150, 150, 40)),
    ("S2 (110X110X25)", (110, 110, 25)),
    ("Massif en G.B (50x50x50)", (50, 50, 50)),
    ("S5(150,0x150,0x40,0)", (150, 150, 40)),
])
def test_m1_direct_semelle(text, expected):
    dims = parse_direct(text)
    assert dims is not None
    L, l, h = dims.interpret("semelle")
    assert (L, l, h) == pytest.approx(expected)


def test_m1_direct_ponctue():
    dims = parse_direct("b=0.5 h=0.4 L=3.2", default_unit="m")
    assert dims is not None
    L, b, h = dims.interpret("longrine")
    assert (L, b, h) == pytest.approx((3.2, 0.5, 0.4))


def test_m1_prioritaire_sur_cotation():
    r = _resolver()
    res = r.resolve_element(ElementSpec(
        label="S4", category="semelle", declared_text="S4(150x150x40)",
        axis_start="4", axis_end="5", pos_start=973, pos_end=1211))
    assert res.method is Method.DIRECT
    assert res.L == pytest.approx(150)
    assert res.surface_m2 == pytest.approx(2.25)


def test_regle_pas_de_surface_sans_x_et_y():
    """'#T8 e=15cm' ne donne qu'une epaisseur : aucune surface inventee."""
    r = _resolver()
    res = r.resolve_element(ElementSpec(
        label="Zone #T8", category="dalle", declared_text="#T8 e=15cm"))
    assert res.surface_m2 is None
    assert res.status is Status.NON_RESOLU


# ---------------------------------------------------------------------------
# 5. M4 : echelle
# ---------------------------------------------------------------------------
def test_calibration_echelle():
    # 248 cm ~ 200 px, 725 cm ~ 580 px, etc. (0.80 px/cm)
    pairs = [(200, 248), (580, 725), (190, 238), (382, 478), (212, 265), (368, 460)]
    model = calibrate_scale(pairs)
    assert model is not None
    assert model.factor == pytest.approx(1.24, abs=0.03)
    assert model.r2 > 0.99
    assert model.n_rejected == 0


def test_calibration_rejette_outlier():
    pairs = [(200, 248), (580, 725), (190, 238), (382, 478), (212, 265),
             (368, 460), (100, 900)]   # 900 cm pour 100 px : aberrant
    model = calibrate_scale(pairs)
    assert model is not None
    assert model.n_rejected == 1
    assert 6 in model.rejected


def test_m4_estimation_sans_cote():
    r = _resolver()
    r.set_scale(calibrate_scale([(200, 248), (580, 725), (190, 238)]))
    res = r.resolve_element(ElementSpec(
        label="S?", category="semelle", bbox=(0, 0, 124, 124)))
    assert res.method is Method.ECHELLE
    assert res.status is Status.ESTIME
    assert res.surface_m2 == pytest.approx((124 * 1.24) ** 2 * 1e-4, rel=0.02)


def test_recoupement_m4_ecart_leve_a_verifier():
    r = _resolver()
    r.set_scale(calibrate_scale([(200, 248), (580, 725), (190, 238)]))
    # declare 150 cm mais bbox 200 px -> ~250 cm mesures
    res = r.resolve_element(ElementSpec(
        label="S4", category="semelle", declared_text="S4(150x150x40)",
        bbox=(0, 0, 200, 124)))
    assert res.secondary is not None
    assert res.secondary.method is Method.ECHELLE
    assert res.secondary.ok is False
    assert res.status is Status.A_VERIFIER
    assert res.secondary.value == pytest.approx(248, abs=2)
    assert res.confidence < 0.90


# ---------------------------------------------------------------------------
# 6. Chaine volontairement fausse
# ---------------------------------------------------------------------------
def test_chaine_fausse_leve_a_verifier():
    r = _resolver()
    r.add_cote(Cote(value=2500, start=0, end=2414, direction="X", label="1-8 FAUSSE"))
    span = r.resolve_span("X", 0, 2414)
    assert span.status is Status.A_VERIFIER
    assert span.value == pytest.approx(2500)      # valeur declaree conservee
    assert span.residual == pytest.approx(86)
    controls = r.closure_controls()
    assert any(not c.ok for c in controls)
    assert any("FERMETURE X" in c.chain and not c.ok for c in controls)


def test_chaine_fausse_element_a_verifier():
    r = _resolver()
    r.add_cote(Cote(value=2500, start=0, end=2414, direction="X", label="1-8 FAUSSE"))
    res = r.resolve_element(ElementSpec(
        label="Emprise", category="emprise", axis_start="1", axis_end="8",
        pos_start=0, pos_end=2414, width=1009))
    assert res.status is Status.A_VERIFIER


# ---------------------------------------------------------------------------
# 7. M5 : polygones
# ---------------------------------------------------------------------------
def test_polygone_avec_trou():
    outer = [(0, 0), (10, 0), (10, 10), (0, 10)]
    hole = [(4, 4), (6, 4), (6, 6), (4, 6)]
    assert polygon_area(outer) == pytest.approx(100)
    assert polygon_area_with_holes([outer, hole]) == pytest.approx(96)


def test_polygone_bulges_cercle():
    """4 arcs de bulge 0.4142 (90 deg) -> aire ~ pi r^2."""
    r = 1.0
    b = math.tan(math.pi / 8)   # 90 deg
    vertices = [(r, 0, b), (0, r, b), (-r, 0, b), (0, -r, b)]
    assert polyline_area(vertices) == pytest.approx(math.pi * r * r, rel=0.01)


def test_surface_polygone_hatch():
    r = _resolver()
    # contour irregulier en cm, echelle 1 -> surface en cm2 puis m2
    outer = [(0, 0), (300, 0), (300, 200), (150, 260), (0, 200)]
    res = r.resolve_element(ElementSpec(
        label="Zone #T8", category="dalle", polygon=[outer]))
    expected_cm2 = polygon_area(outer)
    assert res.surface_m2 == pytest.approx(expected_cm2 * 1e-4, rel=1e-6)
    assert res.method is Method.POLYGONE


# ---------------------------------------------------------------------------
# 8. Overlay et format
# ---------------------------------------------------------------------------
def test_overlay_label_format():
    r = _resolver()
    res = r.resolve_element(ElementSpec(
        label="S4", category="semelle", declared_text="S4(150x150x40)"))
    res.method = Method.SOMME
    res.status = Status.OK
    res.confidence = 1.0
    res.L = 250
    res.l = 170
    res.surface_m2 = 4.25
    assert overlay_label(res) == "S=4,25 m² [SOM OK 1.00]"
    assert format_fr(4.25) == "4,25"


def test_overlay_label_a_verifier():
    r = _resolver()
    res = r.resolve_element(ElementSpec(
        label="S4", category="semelle", declared_text="S4(150x150x40)"))
    res.status = Status.A_VERIFIER
    res.method = Method.SOMME
    res.confidence = 0.87
    res.surface_m2 = 4.25
    assert overlay_label(res) == "S=4,25 m² [SOM A VERIFIER 0.87]"


# ---------------------------------------------------------------------------
# 9. NON_RESOLU
# ---------------------------------------------------------------------------
def test_non_resolu_sans_rien():
    r = _resolver()
    res = r.resolve_element(ElementSpec(label="X1", category="semelle"))
    assert res.status is Status.NON_RESOLU
    assert res.method is Method.AUCUNE
    assert res.surface_m2 is None
    assert res.confidence == 0.0


# ---------------------------------------------------------------------------
# 10. Lecture de cotations par centres de nombres
# ---------------------------------------------------------------------------
def test_cotations_depuis_centres_de_nombres():
    items = [(248, 100, 500), (725, 390, 502), (238, 580, 499),
             (478, 740, 501), (265, 900, 500)]
    cotes = cotations_from_number_centers(items, direction="X")
    assert len(cotes) == 5
    assert [c.value for c in cotes] == [248, 725, 238, 478, 265]
    assert all(c.direction == "X" for c in cotes)
    assert all(c.drawing_length > 0 for c in cotes)


def test_snap_extremites():
    """Deux cotes qui se touchent a 1 cm pres forment une chaine continue."""
    g = AxisGraph("X", [Cote(value=100, start=0, end=100, id="A"),
                        Cote(value=200, start=101, end=301, id="B")], snap_tol=2.0)
    g.build()
    chain = g.chain_between(0, 301)
    assert chain is not None
    assert [c.id for c in chain] == ["A", "B"]


# ---------------------------------------------------------------------------
# 11. Excel
# ---------------------------------------------------------------------------
def _plan_fixture(tmp_path):
    r = _resolver(scale=True)
    specs = [
        ElementSpec(label="S4", category="semelle", declared_text="S4(150x150x40)",
                    element_id="S4-1", axis_start="4", axis_end="5",
                    pos_start=973, pos_end=1211),
        ElementSpec(label="CH axe 6-7", category="chainage",
                    declared_text="CH(20x30)", element_id="CH1",
                    axis_start="6", axis_end="7", pos_start=1689, pos_end=1954,
                    width_start=120, width_end=120),
        ElementSpec(label="Emprise", category="emprise", element_id="EMP",
                    axis_start="1", axis_end="8", pos_start=0, pos_end=2414,
                    width=1009),
        ElementSpec(label="Inconnu", category="semelle", element_id="NON1"),
    ]
    plan = r.resolve_plan(specs)
    path = tmp_path / "metre_dimensions.xlsx"
    plan.to_workbook(path)
    return plan, path


def test_excel_cinq_feuilles(tmp_path):
    import openpyxl

    _plan, path = _plan_fixture(tmp_path)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Elements", "Cotations", "Axes", "Controles", "A_verifier"]


def test_excel_formules_vivantes(tmp_path):
    import openpyxl

    mplan, path = _plan_fixture(tmp_path)
    ws = openpyxl.load_workbook(path)["Elements"]
    rows = {ws.cell(row=r, column=1).value: r for r in range(3, ws.max_row + 1)}

    # chainage : L = cote - demi-largeurs, en formule vivante
    r_ch = rows["CH1"]
    assert str(ws.cell(row=r_ch, column=6).value).startswith("=")
    assert "Cotations!" in str(ws.cell(row=r_ch, column=6).value)
    # chainage : pas de surface "plan", mais un volume L x l x h vivant
    assert ws.cell(row=r_ch, column=9).value is None
    assert str(ws.cell(row=r_ch, column=10).value).startswith("=ROUND(")
    assert f"F{r_ch}" in str(ws.cell(row=r_ch, column=10).value)

    # semelle : surface = L x l en formule vivante
    r_s4 = rows["S4-1"]
    assert ws.cell(row=r_s4, column=9).value is not None
    assert str(ws.cell(row=r_s4, column=9).value).startswith("=ROUND(")

    # emprise : 243,57 m2
    r_emp = rows["EMP"]
    assert ws.cell(row=r_emp, column=9).data_type == "f"
    assert mplan.elements[2].surface_m2 == pytest.approx(243.57, abs=0.01)

    # non resolu : statut et couleur rouge pale
    r_non = rows["NON1"]
    assert ws.cell(row=r_non, column=14).value == "NON_RESOLU"


def test_excel_couleurs_statut(tmp_path):
    import openpyxl

    _plan, path = _plan_fixture(tmp_path)
    ws = openpyxl.load_workbook(path)["Elements"]
    statuts = {ws.cell(row=r, column=14).value: ws.cell(row=r, column=14).fill.fgColor.rgb
               for r in range(3, ws.max_row + 1)}
    assert str(statuts["OK"]).endswith("C6F6D5")
    assert str(statuts["NON_RESOLU"]).endswith("FED7D7")


def test_excel_formule_lisible_en_texte(tmp_path):
    import openpyxl

    _plan, path = _plan_fixture(tmp_path)
    ws = openpyxl.load_workbook(path)["Elements"]
    rows = {ws.cell(row=r, column=1).value: r for r in range(3, ws.max_row + 1)}
    cell = ws.cell(row=rows["CH1"], column=12)
    assert cell.value == "=265-120/2-120/2"
    assert cell.data_type == "s"      # texte, pas une formule Excel


# ---------------------------------------------------------------------------
# 12. Journal / tracabilite
# ---------------------------------------------------------------------------
def test_journal_des_decisions():
    r = _resolver()
    r.resolve_element(ElementSpec(
        label="CH axe 6-7", category="chainage", axis_start="6", axis_end="7",
        pos_start=1689, pos_end=1954, width_start=120, width_end=120))
    assert any("mode net" in line for line in r.log)


def test_unites_metres():
    r = DimensionResolver(ResolverConfig(unit="m"))
    r.add_cotes([Cote(value=2.48, start=0, end=248, direction="X"),
                 Cote(value=7.25, start=248, end=973, direction="X")])
    res = r.resolve_element(ElementSpec(
        label="E", category="emprise", pos_start=0, pos_end=973, width=10.09))
    assert res.L == pytest.approx(9.73)
    assert res.surface_m2 == pytest.approx(98.18, abs=0.02)


# ---------------------------------------------------------------------------
# 13. Overlay : integration du label de metre
# ---------------------------------------------------------------------------
def test_overlay_dessine_label_de_metre(tmp_path):
    cv2 = pytest.importorskip("cv2")
    np = pytest.importorskip("numpy")
    from core.bet_contracts import DrawingElement
    from core.dimension_resolver import (
        Method as M, OverlayAnnotation, ResolvedElement, Status as S,
    )
    from extractors.vision_extractor import PlanBAVisionExtractor

    img = np.full((240, 360, 3), 255, np.uint8)
    elements = [
        DrawingElement(category="semelle", label="S4", bbox=(20, 40, 60, 60)),
        DrawingElement(category="semelle", label="S5", bbox=(150, 40, 60, 60)),
    ]
    res = ResolvedElement(element_id="S4", category="semelle", label="S4",
                          surface_m2=4.25, method=M.SOMME, status=S.OK,
                          confidence=1.0)
    ann = OverlayAnnotation("S=4,25 m² [SOM OK 1.00]", "OK", (0, 180, 0))
    out = tmp_path / "overlay.png"
    path = PlanBAVisionExtractor.render_audit_overlay(
        None, img, elements, out_path=str(out),
        metre_annotations=[ann, None])
    written = cv2.imread(path)
    assert written is not None
    # la bande sous la bbox 1 contient le label de metre (vert)
    band = written[60:100, 15:340]
    green = ((band[:, :, 1] > 120) & (band[:, :, 0] < 100)).sum()
    assert green > 20
    _ = res


# ---------------------------------------------------------------------------
# 14. Pont vision_bridge (non bloquant)
# ---------------------------------------------------------------------------
def test_bridge_dimension_resolution(tmp_path):
    from core.bet_contracts import DrawingElement
    from core.vision_bridge import run_dimension_resolution

    elements = [
        DrawingElement(category="semelle", label="S4", bbox=(10, 10, 40, 40),
                       dimensions={"a": 150, "b": 150, "h": 40}),
        DrawingElement(category="semelle", label="S5", bbox=(60, 10, 40, 40),
                       dimensions={"a": 150, "b": 150, "h": 40}),
    ]
    result = run_dimension_resolution(
        elements=elements, output_dir=str(tmp_path), log=lambda _m: None)
    assert result["xlsx"] is not None
    assert len(result["annotations"]) == 2
    assert result["aligned"] is True
    assert result["plan"].elements[0].method is Method.DIRECT


def test_bridge_ne_leve_jamais(tmp_path):
    from core.vision_bridge import run_dimension_resolution

    result = run_dimension_resolution(output_dir=str(tmp_path),
                                      log=lambda _m: None)
    assert result["xlsx"] is None
    assert result["annotations"] == []
