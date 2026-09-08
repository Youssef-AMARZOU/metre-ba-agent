#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tests/test_pont_gabarit.py -- Pont build_metre.py -> gabarit 5 feuilles.

Verifie : adaptateur plan_data -> plat, resolution du template (dev +
frozen), generation 5 feuilles via MetreGenerator (defaut), allocation
dynamique sans chevauchement, repli historique.
"""
import json
import os
import shutil
import sys
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE_REPO = os.path.join(ROOT, "templates", "modele_metre_BA.xlsx")

PLAN_RICHE = {
    "projet": {"nom": "Projet Pont Test", "date": None},
    "catalogue_types": {
        "semelles": {
            "S1": {"a": 1.0, "b": 1.0, "h": 0.25,
                   "ferr_x": {"nb": 6, "phi": 12},
                   "ferr_y": {"nb": 6, "phi": 12}},
            "S2": {"a": 1.1, "b": 1.1, "h": 0.25,
                   "ferr_x": {"nb": 7, "phi": 12},
                   "ferr_y": {"nb": 7, "phi": 12}},
        },
        "poteaux": {
            "P1": {"a": 0.25, "b": 0.35,
                   "long_bars": [{"nb": 6, "phi": 14}],
                   "cadres": {"phi": 6, "esp": 0.15}},
        },
        "poutres": {
            "N1": {"b": 0.20, "h": 0.30,
                   "filants_inf": [{"nb": 3, "phi": 14}],
                   "filants_sup": [{"nb": 2, "phi": 12}],
                   "cadres": {"phi": 6, "esp": 0.18}},
        },
    },
    "implantations": {
        "semelles": [
            {"id": "S1_1", "type": "S1", "axe": "A", "file": "1"},
            {"id": "S2_1", "type": "S2", "axe": "B", "file": "2"},
        ],
        "poteaux": [
            {"id": "P1", "type": "P1", "axe": "A", "file": "1",
             "hauteur": 3.0},
        ],
        "poutres": [
            {"id": "N1", "type": "N1", "axe": "", "portee": 4.0},
        ],
    },
}


class TestAdaptateur(unittest.TestCase):
    def test_plan_vers_plat(self):
        from build_metre import adapter_plan_vers_injecteur
        payload = adapter_plan_vers_injecteur(PLAN_RICHE)
        self.assertEqual(payload["projet"], "Projet Pont Test")
        self.assertEqual(len(payload["semelles"]), 2)
        s1 = payload["semelles"][0]
        self.assertEqual(
            (s1["type"], s1["axe"], s1["file"], s1["a"], s1["phi"],
             s1["nb_x"], s1["nb_y"]),
            ("S1", "A", "1", 1.0, 12, 6, 6))
        self.assertEqual(len(payload["poteaux"]), 1)
        self.assertEqual(payload["poteaux"][0]["hauteur"], 3.0)
        self.assertEqual(len(payload["poutres"]), 1)
        self.assertEqual(payload["poutres"][0]["portee"], 4.0)

    def test_projet_chaine_conservee(self):
        from build_metre import adapter_plan_vers_injecteur
        payload = adapter_plan_vers_injecteur(
            {"projet": "Nom Direct", "catalogue_types": {},
             "implantations": {}})
        self.assertEqual(payload["projet"], "Nom Direct")


class TestResolutionTemplate(unittest.TestCase):
    def test_dev_trouve_template_repo(self):
        from build_metre import resoudre_chemin_template
        if not os.path.exists(TEMPLATE_REPO):
            self.skipTest("template repo absent")
        if getattr(sys, "frozen", False):
            self.skipTest("mode frozen")
        self.assertEqual(resoudre_chemin_template(), TEMPLATE_REPO)

    def test_frozen_utilise_meipass(self):
        import tempfile
        from build_metre import resoudre_chemin_template
        if not os.path.exists(TEMPLATE_REPO):
            self.skipTest("template repo absent")
        with tempfile.TemporaryDirectory() as td:
            dest = os.path.join(td, "templates")
            os.makedirs(dest)
            shutil.copy(TEMPLATE_REPO,
                        os.path.join(dest, "modele_metre_BA.xlsx"))
            old_frozen = getattr(sys, "frozen", None)
            old_meipass = getattr(sys, "_MEIPASS", None)
            sys.frozen = True
            sys._MEIPASS = td
            try:
                # Sans base_dir : la branche frozen s'applique
                trouve = resoudre_chemin_template()
                self.assertEqual(
                    trouve,
                    os.path.join(dest, "modele_metre_BA.xlsx"))
            finally:
                if old_frozen is None:
                    del sys.frozen
                else:
                    sys.frozen = old_frozen
                if old_meipass is None:
                    if hasattr(sys, "_MEIPASS"):
                        del sys._MEIPASS
                else:
                    sys._MEIPASS = old_meipass

    def test_regeneration_si_absent(self):
        import tempfile
        from build_metre import resoudre_chemin_template
        with tempfile.TemporaryDirectory() as td:
            trouve = resoudre_chemin_template(base_dir=td)
            self.assertTrue(os.path.exists(trouve))
            wb = openpyxl.load_workbook(trouve)
            self.assertIn("05_Catalogue_Armatures_Standard", wb.sheetnames)


class TestGenererGabarit(unittest.TestCase):
    def setUp(self):
        import tempfile
        self._td = tempfile.TemporaryDirectory()
        self.addCleanup(self._td.cleanup)
        self.out = os.path.join(self._td.name, "metre.xlsx")

    def _generer(self, plan):
        from build_metre import MetreGenerator
        gen = MetreGenerator(plan)  # moteur gabarit par defaut
        return gen.generer(self.out)

    def test_cinq_feuilles(self):
        if not os.path.exists(TEMPLATE_REPO):
            self.skipTest("template repo absent")
        self._generer(PLAN_RICHE)
        wb = openpyxl.load_workbook(self.out)
        self.assertEqual(wb.sheetnames,
                         ["01_Detail_Quantitatif", "02_Armatures",
                          "03_Attachement_Ferraillage", "04_GO_Attachement",
                          "05_Catalogue_Armatures_Standard"])

    def test_donnees_reelles_injectees(self):
        if not os.path.exists(TEMPLATE_REPO):
            self.skipTest("template repo absent")
        self._generer(PLAN_RICHE)
        wb = openpyxl.load_workbook(self.out)
        ws1 = wb["01_Detail_Quantitatif"]
        textes = " | ".join(str(ws1.cell(r, 4).value or "")
                            for r in range(1, ws1.max_row + 1))
        self.assertIn("Fouille Semelle S1", textes)
        self.assertIn("Semelle S2", textes)
        self.assertIn("Fût P1", textes)
        ws2 = wb["02_Armatures"]
        textes2 = " | ".join(str(ws2.cell(r, 1).value or "")
                             for r in range(1, ws2.max_row + 1))
        self.assertIn("S1", textes2)
        self.assertIn("Armature INF X", textes2)
        self.assertIn("ARM LONG", textes2)

    def test_liens_feuille4_dynamiques(self):
        if not os.path.exists(TEMPLATE_REPO):
            self.skipTest("template repo absent")
        self._generer(PLAN_RICHE)
        wb = openpyxl.load_workbook(self.out)
        ws4 = wb["04_GO_Attachement"]
        e4 = ws4.cell(4, 5).value or ""
        e7 = ws4.cell(7, 5).value or ""
        self.assertIn("01_Detail_Quantitatif", e4)
        self.assertIn("02_Armatures", e7)
        # Les liens doivent pointer sur des sous-totaux existants (formules)
        import re
        ws1 = wb["01_Detail_Quantitatif"]
        m = re.search(r"!K(\d+)$", e4)
        self.assertIsNotNone(m, f"lien inattendu : {e4}")
        cible = int(m.group(1))
        self.assertTrue(str(ws1.cell(cible, 11).value).startswith("=SUM("))

    def test_allocation_dynamique_10_semelles(self):
        """10 semelles : sous-totaux apres les entrees, plages SUM justes."""
        if not os.path.exists(TEMPLATE_REPO):
            self.skipTest("template repo absent")
        import copy
        plan = copy.deepcopy(PLAN_RICHE)
        plan["implantations"]["semelles"] = [
            {"id": f"S1_{i}", "type": "S1",
             "axe": "A", "file": str(i)} for i in range(1, 11)]
        self._generer(plan)
        wb = openpyxl.load_workbook(self.out)
        ws1 = wb["01_Detail_Quantitatif"]
        # Premier sous-total : 10 entrees J9..J18
        subs = [r for r in range(1, ws1.max_row + 1)
                if ws1.cell(r, 1).value == "Ss-total"]
        self.assertGreaterEqual(len(subs), 3)
        self.assertEqual(ws1.cell(subs[0], 10).value, "=SUM(J9:J18)")
        # Feuille 2 : 10 semelles x 3 lignes (mere+X+Y) + poteau (3 lignes)
        # + poutre (4 lignes) => lignes 4..40
        ws2 = wb["02_Armatures"]
        r_long = next(r for r in range(1, ws2.max_row + 1)
                      if ws2.cell(r, 1).value == "LONGUEUR TOTALE (ml)")
        self.assertEqual(ws2.cell(r_long, 11).value, "=SUM(K4:K40)")

    def test_moteur_historique_2_feuilles(self):
        from build_metre import MetreGenerator
        gen = MetreGenerator(PLAN_RICHE, moteur="historique")
        gen.generer(self.out)
        wb = openpyxl.load_workbook(self.out)
        self.assertEqual(len(wb.sheetnames), 2)


if __name__ == "__main__":
    unittest.main()
