# -*- coding: utf-8 -*-
"""tests/test_surface_metre.py -- Metre des surfaces (plans d'architecture)."""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

cv2 = pytest.importorskip("cv2")
np = pytest.importorskip("numpy")

from extractors.surface_extractor import (  # noqa: E402
    PlanSurfaceExtractor, RoomSurface, SurfaceReport, is_dimension_token,
    parse_length, parse_room_dims)
from core.surface_metre_excel import SURFACE_SHEETS, SurfaceMetreExcel  # noqa: E402
from core.vision_bridge import plan_data_is_placeholder, run_surface_metre  # noqa: E402


# ----------------------------------------------------------------------------
# Unites
# ----------------------------------------------------------------------------

@pytest.mark.parametrize("text,expected", [
    ("16'-0\"", 4.877), ("16-0\"", 4.877), ("13-6\"", 4.115), ("12.3'", 3.749),
    ("3.50", 3.5), ("3,50 m", 3.5), ("350", 3.5), ("350 cm", 3.5),
    ("3500 mm", 3.5), ("4200", 4.2), ("0.90", 0.9), ("12\"", 0.305),
])
def test_parse_length(text, expected):
    assert parse_length(text) == pytest.approx(expected, abs=0.002)


def test_parse_length_invalide():
    assert parse_length("KITCHEN") is None
    assert parse_length("") is None
    assert parse_length("P12") is None


@pytest.mark.parametrize("text,expected", [
    ("18'x10\"", (5.486, 3.048)),      # OCR ' lu " : contexte piece -> pieds
    ("13'×10'", (3.962, 3.048)),
    ("4.20x3.50", (4.2, 3.5)),
    ("420 x 350", (4.2, 3.5)),
    ("12'x.12\"", (3.658, 3.658)),
])
def test_parse_room_dims(text, expected):
    got = parse_room_dims(text)
    assert got is not None
    assert got[0] == pytest.approx(expected[0], abs=0.002)
    assert got[1] == pytest.approx(expected[1], abs=0.002)


def test_room_dims_rejette_valeurs_absurdes():
    assert parse_room_dims("P1") is None
    assert parse_room_dims("0.2x0.3") is None     # < 0.5 m : pas une piece
    assert parse_room_dims("5000x4000") == pytest.approx((5.0, 4.0))


def test_is_dimension_token():
    assert is_dimension_token("57-0\"")
    assert is_dimension_token("3.50 m")
    assert is_dimension_token("4200")
    assert not is_dimension_token("18'x10'")
    assert not is_dimension_token("KITCHEN")
    assert not is_dimension_token("12")


# ----------------------------------------------------------------------------
# Plan synthetique : 2 pieces, echelle 100 px/m
# ----------------------------------------------------------------------------

def _synthetic_plan():
    """Batiment 8.00 x 5.00 m a 100 px/m : piece A 5x5, piece B 3x5, porte 0.9 m."""
    S = 100
    img = np.full((800, 1100, 3), 255, np.uint8)
    x0, y0 = 150, 150
    W, H = 8 * S, 5 * S
    cv2.rectangle(img, (x0, y0), (x0 + W, y0 + H), (0, 0, 0), 8)
    xm = x0 + 5 * S
    cv2.line(img, (xm, y0), (xm, y0 + 2 * S), (0, 0, 0), 8)            # cloison
    cv2.line(img, (xm, y0 + 2 * S + 90), (xm, y0 + H), (0, 0, 0), 8)   # porte 0.9 m
    # Mobilier fin dans la piece A (ne doit pas devenir une piece)
    cv2.rectangle(img, (x0 + 60, y0 + 300), (x0 + 260, y0 + 450), (0, 0, 0), 1)
    ocr = [
        {"text": "SALON", "bbox": (x0 + 120, y0 + 330, 80, 20), "conf": 0.98},   # dans le lit
        {"text": "CHAMBRE 1", "bbox": (xm + 80, y0 + 220, 110, 20), "conf": 0.97},
        {"text": "RDC", "bbox": (20, 20, 40, 18), "conf": 0.9},
        # Chaine de cotes horizontale au-dessus (centres aux milieux)
        {"text": "5.00", "bbox": (x0 + 250 - 20, y0 - 60, 40, 18), "conf": 0.95},
        {"text": "3.00", "bbox": (xm + 150 - 20, y0 - 60, 40, 18), "conf": 0.95},
        {"text": "8.00", "bbox": (x0 + 400 - 20, y0 - 110, 40, 18), "conf": 0.95},
        # Chaine verticale a gauche
        {"text": "5.00", "bbox": (x0 - 70, y0 + 250 - 20, 18, 40), "conf": 0.95},
    ]
    return img, ocr


def test_segmentation_synthetique():
    img, ocr = _synthetic_plan()
    rep = PlanSurfaceExtractor().extract_from_image(img, ocr_items=ocr)
    assert rep.level == "RDC"
    assert rep.scale_px_per_m == pytest.approx(100.0, rel=0.08)
    assert rep.scale_confidence >= 0.6
    by = {r.label: r for r in rep.rooms}
    assert "SALON" in by and "CHAMBRE 1" in by
    assert by["SALON"].source == "segmentation"
    assert by["SALON"].area_m2 == pytest.approx(25.0, rel=0.12)
    assert by["CHAMBRE 1"].area_m2 == pytest.approx(15.0, rel=0.12)
    assert by["SALON"].perimeter_m == pytest.approx(20.0, rel=0.15)
    assert rep.total_area_m2 == pytest.approx(40.0, rel=0.12)
    assert not any(w.startswith("[ALERTE") for w in rep.warnings)


def test_cote_explicite_prioritaire_et_calibration_manuelle(tmp_path):
    img, ocr = _synthetic_plan()
    ocr = ocr + [{"text": "5.00x5.00", "bbox": (150 + 180, 150 + 60, 90, 18), "conf": 0.9}]
    rep = PlanSurfaceExtractor().extract_from_image(
        img, ocr_items=ocr, known_width_m=8.0, overlay_dir=str(tmp_path))
    by = {r.label: r for r in rep.rooms}
    assert by["SALON"].source == "cote_explicite"
    assert by["SALON"].area_m2 == pytest.approx(25.0)
    assert rep.scale_confidence == 0.95
    assert rep.scale_estimates[0]["method"] == "manuel_hors_tout"
    assert Path(rep.overlay_path).exists()


def test_alerte_cotes_hors_tout_contradictoires():
    img, ocr = _synthetic_plan()
    # Deux cotes hors-tout centrees, lectures concurrentes de la MEME cote
    # (8.00 m a 100 px/m). Ecart plausible -> alerte.
    ocr = ocr + [{"text": "9.50", "bbox": (150 + 400 - 20, 150 - 110, 40, 18),
                  "conf": 0.9}]
    rep = PlanSurfaceExtractor().extract_from_image(img, ocr_items=ocr)
    assert any("contradictoires" in w for w in rep.warnings)
    assert rep.scale_px_per_m == pytest.approx(100.0, rel=0.05)


def test_cote_hors_tout_erronee_ecartee_par_les_chaines():
    """Une cote hors-tout aberrante (24.00 pour une emprise de 8.00 m) ne doit
    pas fausser l'echelle : les chaines arbitrent, l'ecart est trace."""
    img, ocr = _synthetic_plan()
    ocr = ocr + [{"text": "24.00", "bbox": (150 + 400 - 20, 150 - 140, 40, 18),
                  "conf": 0.9}]
    rep = PlanSurfaceExtractor().extract_from_image(img, ocr_items=ocr)
    assert rep.scale_px_per_m == pytest.approx(100.0, rel=0.05)
    assert any("incoherente avec les chaines" in w for w in rep.warnings)
    assert any(e["method"] == "hors_tout_x_ignoree" for e in rep.scale_estimates)


def test_sans_texte_sans_hallucination():
    img = np.full((400, 400, 3), 255, np.uint8)
    rep = PlanSurfaceExtractor().extract_from_image(img, ocr_items=[])
    assert rep.rooms == [] and rep.total_area_m2 == 0.0
    assert rep.warnings


# ----------------------------------------------------------------------------
# Excel + pont
# ----------------------------------------------------------------------------

def test_excel_surfaces(tmp_path):
    from openpyxl import load_workbook
    rep = SurfaceReport(level="RDC", scale_px_per_m=100.0, scale_confidence=0.9)
    rep.rooms = [
        RoomSurface("SALON", "RDC", 5.0, 4.0, 20.0, 19.5, 18.0, "cote_explicite", 0.9),
        RoomSurface("CHAMBRE 1", "RDC", None, None, None, 12.0, 14.0, "segmentation", 0.7),
        RoomSurface("WC", "RDC", source="aucune", confidence=0.0, warnings=["Surface NON TROUVEE"]),
    ]
    rep2 = SurfaceReport(level="ETAGE 1", page_num=2, scale_px_per_m=100.0, scale_confidence=0.8)
    rep2.rooms = [RoomSurface("CHAMBRE 2", "ETAGE 1", 4.0, 3.0, 12.0, None, 14.0, "cote_explicite", 0.9)]
    x = SurfaceMetreExcel("Test", hauteur_sous_plafond_m=2.80)
    out = x.export([rep, rep2], tmp_path / "s.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == list(SURFACE_SHEETS)
    ws = wb[SURFACE_SHEETS[0]]
    assert ws["F6"].value == '=IF(AND(ISNUMBER(D6),ISNUMBER(E6)), ROUND(D6*E6, 2), "")'
    assert ws["H6"].value == "=IF(ISNUMBER(F6), F6, IF(ISNUMBER(G6), G6, 0))"
    lay = x.layout["pieces"]
    assert ws[f"H{lay['total_row']}"].value == f"=ROUND(SUM(H{lay['start']}:H{lay['end']}), 2)"
    assert "SUMIF" in ws[f"H{lay['sub_start']}"].value
    ws2 = wb[SURFACE_SHEETS[1]]
    assert ws2["B4"].value == 2.80
    assert "$B$4" in ws2["D8"].value   # enduit = perimetre x HSP
    assert ws2["F6"].value == "=ROUND(D6*E6, 2)"


def test_run_surface_metre_image(tmp_path):
    img, ocr = _synthetic_plan()
    # Injecte les tokens via un routeur OCR factice
    class FakeOCR:
        engine_name = "fake"
        available = True
        def recognize(self, img_bgr, offset=(0, 0)):
            return ocr
    import extractors.surface_extractor as se
    orig = se.PlanSurfaceExtractor.ocr
    se.PlanSurfaceExtractor.ocr = property(lambda self: FakeOCR())
    try:
        path = tmp_path / "plan.png"
        cv2.imwrite(str(path), img)
        res = run_surface_metre(str(path), str(tmp_path), "Synth", log=lambda m: None)
    finally:
        se.PlanSurfaceExtractor.ocr = orig
    assert res["xlsx"] and Path(res["xlsx"]).exists()
    assert res["total_m2"] == pytest.approx(40.0, rel=0.12)
    assert res["summary"]["n_rooms"] == 2 and res["summary"]["by_level"]["RDC"] > 0


def test_plan_data_placeholder():
    from core.local_extractor import partial_plan_data
    assert plan_data_is_placeholder(partial_plan_data("x"))
    assert not plan_data_is_placeholder({"catalogue_types": {"semelles": {"S1": {"a": 1}}}})
