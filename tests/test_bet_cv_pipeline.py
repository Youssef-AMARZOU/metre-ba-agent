# -*- coding: utf-8 -*-
"""
tests/test_bet_cv_pipeline.py -- Verification load-bearing du pipeline
Computer Vision + MoroccanBETEngine + classeur 5 feuilles.

Criteres d'acceptation :
  1. Auto-correction des unites (300x300 mm -> 0.30 m ; 140x140 cm -> 1.40 m)
  2. Portee entre nus (4.80 axe-axe, poteaux 0.30 -> 4.50)
  3. HA18 invalide + ecart de quantite Coffrage vs Nomenclature
  4. Classeur .xlsx 5 feuilles sans erreur de formule (recalcul Excel reel
     si disponible, sinon validation syntaxique + references)
  + couches CV (SHX/scan, tables, contours, cKDTree, anti-hallucination,
    bulles d'axes, overlay, routeur VLM mocke, pont legacy).
"""
from __future__ import annotations

import math
import os
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.bet_contracts import (  # noqa: E402
    DrawingElement, ExtractedPlanRecord, records_to_legacy_plan_data)
from core.moroccan_bet_engine import (  # noqa: E402
    MOROCCAN_REBAR_CATALOG, MOROCCAN_STEEL_RATIO_BOUNDS,
    MoroccanBETEngine, PlanBAGatekeeper)
from core.populate_modele import (  # noqa: E402
    PlanBAExcelPopulator, SHEET_NAMES, injecter_records_certifies)
from core.populate_modele_bet import BPU_ITEMS, optimize_12m_stock_bars  # noqa: E402

cv2 = pytest.importorskip("cv2")
np = pytest.importorskip("numpy")
fitz = pytest.importorskip("fitz")

from extractors.vision_extractor import PlanBAVisionExtractor  # noqa: E402
from extractors.vlm_router import VLMRouter  # noqa: E402


# ============================================================================
# Fixtures
# ============================================================================

@pytest.fixture
def engine():
    return MoroccanBETEngine()


@pytest.fixture
def gate(engine):
    return PlanBAGatekeeper(engine)


def _sample_records():
    """Jeu certifie : S1 en cm, P1 en mm, LG1 mixte, P2 avec HA18 + ecart."""
    return [
        ExtractedPlanRecord(
            category="semelle", label="S1", level="Fondations",
            grid_axes=["A-1", "A-2", "B-1", "B-2"],
            qty_plan_coffrage=4, qty_table_schedule=4,
            raw_dimensions={"A": 140, "B": 140, "h_patin": 35},
            rebar_specs=["12HA14 L=1.30", "12HA14 L=1.30"]),
        ExtractedPlanRecord(
            category="longrine", label="LG1", level="Fondations",
            grid_axes=["A-1..A-2", "B-1..B-2"],
            qty_plan_coffrage=2, qty_table_schedule=2,
            raw_dimensions={"L": 4.80, "b": 25, "h": 45},
            support_left_m=0.30, support_right_m=0.30,
            rebar_specs=["6HA14", "Cad HA8 e=15"]),
        ExtractedPlanRecord(
            category="poteau", label="P1", level="RDC",
            grid_axes=["A-1", "A-2", "B-1", "B-2"],
            qty_plan_coffrage=4, qty_table_schedule=4,
            raw_dimensions={"L": 3.20, "b": 300, "h": 300},
            rebar_specs=["8HA14", "Cad HA8 e=12"]),
        ExtractedPlanRecord(
            category="poteau", label="P2", level="RDC",
            grid_axes=["C-1", "C-2"],
            qty_plan_coffrage=2, qty_table_schedule=3,
            raw_dimensions={"L": 3.20, "b": 25, "h": 40},
            rebar_specs=["6HA18", "Cad HA8 e=15"]),
    ]


# ============================================================================
# 1. Porte 1 : unites
# ============================================================================

def test_porte1_normalisation_mm_cm_vers_m(gate):
    d = gate.normalize_dimensions_to_meters("P1", "poteau", {"b": 300, "h": 300})
    assert d == {"b": 0.30, "h": 0.30}
    d = gate.normalize_dimensions_to_meters("S1", "semelle", {"A": 140, "B": 140, "H": 40})
    assert d == {"A": 1.40, "B": 1.40, "H": 0.40}
    d = gate.normalize_dimensions_to_meters("P1", "poteau", {"b": 25, "h": 40})
    assert d == {"b": 0.25, "h": 0.40}
    d = gate.normalize_dimensions_to_meters("N1", "poutre", {"L": 4500, "b": 0.25, "h": 0.45})
    assert d == {"L": 4.50, "b": 0.25, "h": 0.45}
    # Deja en metres : inchange
    d = gate.normalize_dimensions_to_meters("S2", "semelle", {"A": 1.2, "B": 1.2, "H": 0.3})
    assert d == {"A": 1.2, "B": 1.2, "H": 0.3}


# ============================================================================
# 2. Portee entre nus
# ============================================================================

def test_portee_entre_nus_exacte(engine):
    line = engine.compute_poutre_entre_nus(
        "N1", qty=1, portee_axe_axe=4.80, largeur_poteau_gauche=0.30,
        largeur_poteau_droite=0.30, b_poutre=0.25, h_totale=0.45)
    assert line.longueur == pytest.approx(4.50)
    assert line.hauteur == pytest.approx(0.45)
    assert line.quantite_partielle == pytest.approx(round(4.50 * 0.25 * 0.45, 3))
    assert line.prix_code == "B.06"
    assert not engine.anomalies


def test_portee_dalle_pleine_retombee(engine):
    line = engine.compute_poutre_entre_nus(
        "N2", qty=2, portee_axe_axe=5.00, largeur_poteau_gauche=0.40,
        largeur_poteau_droite=0.20, b_poutre=0.25, h_totale=0.60,
        type_plancher="dalle_pleine", epaisseur_dalle=0.15)
    assert line.longueur == pytest.approx(4.70)
    assert line.hauteur == pytest.approx(0.45)
    assert line.quantite_partielle == pytest.approx(round(2 * 4.70 * 0.25 * 0.45, 3))


def test_portee_negative_bloquante(engine):
    engine.compute_poutre_entre_nus("N9", 1, 0.40, 0.50, 0.50, 0.25, 0.45)
    assert any(a.rule_code == "GEOM_SPAN_NEGATIVE" for a in engine.anomalies)


# ============================================================================
# 3. Diametre invalide + ecart de quantite
# ============================================================================

def test_diametre_invalide_ha18(engine):
    mass = engine.validate_and_get_unit_mass("P2", 18)
    anoms = [a for a in engine.anomalies if a.rule_code == "INVALID_DIAMETER_OCR"]
    assert len(anoms) == 1 and anoms[0].severity == "BLOQUANT"
    assert "HA16" in anoms[0].suggested_fix
    assert mass == MOROCCAN_REBAR_CATALOG[16]
    assert engine.validate_and_get_unit_mass("P2", 15) == MOROCCAN_REBAR_CATALOG[14]
    assert engine.validate_and_get_unit_mass("P2", 12) == 0.888


def test_ecart_quantite_coffrage_vs_nomenclature():
    pop = PlanBAExcelPopulator("Test")
    pop.compute(_sample_records())
    codes = {(a.element_label, a.rule_code) for a in pop.engine.anomalies}
    assert ("P2", "CROSS_CHECK_QTY_MISMATCH") in codes
    assert ("P2", "INVALID_DIAMETER_OCR") in codes
    # S1, LG1 et P1 sont sains : pas d'ecart ni de diametre invalide
    for lbl in ("S1", "LG1", "P1"):
        assert (lbl, "CROSS_CHECK_QTY_MISMATCH") not in codes
        assert (lbl, "INVALID_DIAMETER_OCR") not in codes


# ============================================================================
# Formules acier RPS 2011
# ============================================================================

def test_semelle_proprete_et_glacis(engine):
    bp, ba = engine.compute_semelle_isolee("S1", 1, A=1.40, B=1.40,
                                           h_patin=0.35, h_glacis=0.0)
    assert bp.longueur == pytest.approx(1.60) and bp.largeur == pytest.approx(1.60)
    assert bp.quantite_partielle == pytest.approx(round(1.6 * 1.6 * 0.10, 3))
    assert ba.quantite_partielle == pytest.approx(round(1.4 * 1.4 * 0.35, 3))
    _, ba2 = engine.compute_semelle_isolee("S2", 1, A=2.0, B=2.0, h_patin=0.30,
                                           h_glacis=0.30, a_fut=0.40, b_fut=0.40)
    tronc = 0.30 / 6 * (4.0 + 0.16 + 2.4 * 2.4)
    assert ba2.quantite_partielle == pytest.approx(round(1.2 + tronc, 3))
    # Le temoin F*G*H*I (formule Excel) reproduit le volume reel
    assert ba2.nombre * ba2.longueur * ba2.largeur * ba2.hauteur == pytest.approx(
        ba2.quantite_partielle, abs=0.005)


def test_cadres_et_recouvrements(engine):
    # Cadre 25x45, HA8 : 2(0.20)+2(0.40)+2*0.08 = 1.36 m
    assert engine.stirrup_developed_length(0.25, 0.45, 8) == pytest.approx(1.36)
    assert engine.stirrup_count(4.50, 0.15) == 31
    # Barre HA14 de 13 m : crosses 2*15*0.014 = 0.42 -> 13.42 > 12 -> 1 recouvrement 50*0.014
    assert engine.developed_length_longitudinal(14, 13.0) == pytest.approx(13.42 + 0.70)
    assert engine.developed_length_longitudinal(14, 4.0) == pytest.approx(4.42)
    assert engine.developed_length_longitudinal(14, 4.0, explicit_length=True) == pytest.approx(4.0)


def test_porte3_decodage_armatures(gate):
    rec = ExtractedPlanRecord(category="longrine", label="LG1",
                              raw_dimensions={"L": 4.5, "b": 0.25, "h": 0.45},
                              rebar_specs=["6HA14", "Cad HA8 e=15", "3T12 L=4.20"])
    rows, kg = gate.parse_rebar_specs_to_schedule(rec, rec.raw_dimensions, 2)
    marks = [r.bar_mark for r in rows]
    assert any("6HA14" in m for m in marks)
    assert any("Cad HA8 e=15" in m for m in marks)
    assert any("3HA12" in m for m in marks)
    row12 = next(r for r in rows if r.diameter_mm == 12)
    assert row12.developed_length_m == pytest.approx(4.20)
    assert kg > 0 and not gate.engine.anomalies


def test_ratio_rps_hors_plage(engine):
    from core.bet_contracts import AvantMetreLine
    line = AvantMetreLine("B.05", "poteau", "Poteau P9", "m3", 1, 3.0, 0.3, 0.3,
                          0.27, acier_kg=200.0, label="P9")
    assert not engine.verify_element_consistency(line, None, None)
    assert line.ratio_acier_kg_m3 == pytest.approx(740.7)
    assert any(a.rule_code == "RPS2011_RATIO_OUT_OF_BOUNDS" for a in engine.anomalies)


def test_ffd_barres_12m():
    from core.bet_contracts import RebarScheduleRow
    rows = [RebarScheduleRow("N1", "poutre", "Pos.1", 14, 1, 4, 5.0),
            RebarScheduleRow("N1", "poutre", "Pos.2", 14, 1, 2, 7.0)]
    res = optimize_12m_stock_bars(rows)
    # 4x5 + 2x7 = 34 ml -> 5+7 (x2) + 5+5 (x1) = 3 barres de 12 m
    assert res[14]["nb_stock_bars_12m"] == 3
    assert res[14]["useful_length_m"] == pytest.approx(34.0)
    assert res[14]["waste_length_m"] == pytest.approx(2.0)


# ============================================================================
# 4. Classeur 5 feuilles
# ============================================================================

def _formula_refs_valid(wb) -> list:
    """Validation syntaxique (tokenizer openpyxl) + feuilles referencees."""
    from openpyxl.formula import Tokenizer
    errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                try:
                    tok = Tokenizer(v)
                except Exception as e:
                    errors.append(f"{ws.title}!{cell.coordinate}: {e}")
                    continue
                for t in tok.items:
                    if t.type == "OPERAND" and t.subtype == "RANGE" and "!" in t.value:
                        sheet = t.value.split("!")[0].strip("'")
                        if sheet not in wb.sheetnames:
                            errors.append(f"{ws.title}!{cell.coordinate}: feuille {sheet}")
    return errors


_EXCEL_RECALC_SCRIPT = r"""
import json, sys
import win32com.client
path = sys.argv[1]
excel = win32com.client.DispatchEx("Excel.Application")
excel.Visible = False
excel.DisplayAlerts = False
try:
    wb = excel.Workbooks.Open(path, UpdateLinks=0, ReadOnly=True)
    excel.CalculateFullRebuild()
    data = {}
    for ws in wb.Worksheets:
        v = ws.UsedRange.Value
        if not isinstance(v, tuple):
            v = ((v,),)
        data[ws.Name] = [[(c if isinstance(c, (int, float, str)) or c is None else str(c))
                          for c in row] for row in v]
    wb.Close(False)
finally:
    excel.Quit()
print(json.dumps(data))
"""


def _excel_recalc(path: Path):
    """Recalcul reel via Excel COM dans un sous-processus isole (les
    deconnexions RPC a la fermeture d'Excel ne touchent pas pytest).
    Retourne {feuille: [[valeurs]]} ou None si indisponible."""
    import json
    import subprocess
    try:
        import win32com.client  # noqa: F401
    except Exception:
        return None
    try:
        proc = subprocess.run(
            [sys.executable, "-c", _EXCEL_RECALC_SCRIPT, str(path.resolve())],
            capture_output=True, text=True, timeout=120)
    except Exception:
        return None
    if proc.returncode != 0 or not proc.stdout.strip():
        return None
    try:
        return json.loads(proc.stdout.strip().splitlines()[-1])
    except Exception:
        return None


EXCEL_ERROR_CODES = {-2146826281, -2146826273, -2146826259, -2146826265,
                     -2146826246, -2146826252, -2146826288}


def test_generation_classeur_5_feuilles(tmp_path):
    from openpyxl import load_workbook
    pop = PlanBAExcelPopulator("Construction R+2 - Lot Gros Oeuvre")
    out = pop.process_and_export(_sample_records(), tmp_path / "metre_bet.xlsx")
    assert out.exists()

    wb = load_workbook(out)
    assert wb.sheetnames == list(SHEET_NAMES)
    assert _formula_refs_valid(wb) == []

    ws1 = wb[SHEET_NAMES[0]]
    lay = pop.layout["avant_metre"]
    assert ws1[f"J{lay['start']}"].value == f"=ROUND(F{lay['start']}*G{lay['start']}*H{lay['start']}*I{lay['start']}, 3)"
    assert ws1[f"J{lay['total_row']}"].value == f"=ROUND(SUM(J{lay['start']}:J{lay['end']}), 3)"
    ws2 = wb[SHEET_NAMES[1]]
    a = pop.layout["acier"]
    assert ws2[f"H{a['start']}"].value == f"=ROUND(E{a['start']}*F{a['start']}*G{a['start']}, 2)"
    assert ws2[f"I{a['start']}"].value.startswith(f"=IF($D{a['start']}=6, ROUND($H{a['start']}*0.222")
    ws3 = wb[SHEET_NAMES[2]]
    assert "SUMIF('1_Avant_Metre_BA'!$A$6:$A$500" in ws3["D6"].value
    ws4 = wb[SHEET_NAMES[3]]
    assert ws4["D6"].value == "=IF(B6=0, 0, ROUND(C6/B6, 1))"
    assert 'IF(AND(D6>=E6, D6<=F6), "CONFORME BET MAROC"' in ws4["G6"].value
    assert ws4["A6"].value == "semelle" and ws4["E6"].value == 40.0 and ws4["F6"].value == 95.0
    ws5 = wb[SHEET_NAMES[4]]
    assert ws5["D6"].value == "=B6*12" and ws5["E6"].value == "=ROUND(D6-C6, 2)"

    # Les anomalies attendues figurent dans le journal
    journal = [c.value for row in ws4.iter_rows(min_col=3, max_col=3) for c in row]
    assert "INVALID_DIAMETER_OCR" in journal and "CROSS_CHECK_QTY_MISMATCH" in journal


def test_recalcul_excel_reel_sans_erreur(tmp_path):
    """Load-bearing : Excel recalcule, aucune cellule en erreur, et le total
    beton / acier Excel == temoins Python du moteur deterministe."""
    pop = PlanBAExcelPopulator("Recalcul")
    out = pop.process_and_export(_sample_records(), tmp_path / "recalc.xlsx")
    data = _excel_recalc(out)
    if data is None:
        pytest.skip("Excel COM indisponible : recalcul reel non verifiable")

    for name, grid in data.items():
        for r in grid:
            for v in r:
                assert not (isinstance(v, (int, float)) and int(v) in EXCEL_ERROR_CODES), \
                    f"Erreur de formule Excel dans {name} : {v}"

    lay = pop.layout["avant_metre"]
    grid1 = data[SHEET_NAMES[0]]
    total_xl = grid1[lay["total_row"] - 1][9]      # colonne J (0-based 9)
    total_py = round(sum(l.quantite_partielle for l in pop.avant_metre_lines), 3)
    assert total_xl == pytest.approx(total_py, abs=0.01)
    for i, line in enumerate(pop.avant_metre_lines):
        assert grid1[lay["start"] - 1 + i][9] == pytest.approx(line.quantite_partielle, abs=0.005)

    a = pop.layout["acier"]
    grid2 = data[SHEET_NAMES[1]]
    acier_xl = grid2[a["grand_row"] - 1][a["first_dia_col"] - 1]
    acier_py = 0.0
    for r in pop.rebar_schedule_rows:
        d = r.diameter_mm if r.diameter_mm in MOROCCAN_REBAR_CATALOG else None
        if d in pop.DIAMETERS_COLS:
            acier_py += round(r.total_length_m * MOROCCAN_REBAR_CATALOG[d], 2)
    assert acier_xl == pytest.approx(acier_py, abs=0.05)

    grid3 = data[SHEET_NAMES[2]]
    d3 = pop.layout["dqe"]
    ht = grid3[d3["t_ht"] - 1][5]
    tva = grid3[d3["t_tva"] - 1][5]
    ttc = grid3[d3["t_ttc"] - 1][5]
    assert tva == pytest.approx(round(ht * 0.20, 2), abs=0.01)
    assert ttc == pytest.approx(round(ht + tva, 2), abs=0.01)

    grid4 = data[SHEET_NAMES[3]]
    statuts = {grid4[i][0]: grid4[i][6] for i in range(5, 5 + len(pop.CONTROL_CATEGORIES))}
    assert statuts["poutre"] == "NON APPLICABLE"
    assert statuts["longrine"] in ("CONFORME BET MAROC", "ALERTE HORS PLAGE RPS 2011")


def test_injecter_records_certifies_wrapper(tmp_path):
    out = injecter_records_certifies(_sample_records(), str(tmp_path / "w.xlsx"), "Wrap")
    assert Path(out).exists()


# ============================================================================
# Computer Vision — couche 1 : bascule SHX / scan
# ============================================================================

def _pdf_page(tmp_path, name, n_words, n_lines):
    doc = fitz.open()
    page = doc.new_page(width=842, height=595)
    for i in range(n_words):
        page.insert_text((40 + (i % 10) * 75, 60 + (i // 10) * 20), f"P{i}")
    for i in range(n_lines):
        x = 50 + (i % 700)
        page.draw_line((x, 300 + (i // 700) * 3), (x + 1, 302 + (i // 700) * 3))
    path = tmp_path / name
    doc.save(str(path))
    doc.close()
    return path


def test_needs_vision_fallback_scan_et_shx(tmp_path):
    p_scan = _pdf_page(tmp_path, "scan.pdf", n_words=3, n_lines=10)
    p_vec = _pdf_page(tmp_path, "vector.pdf", n_words=60, n_lines=100)
    p_shx = _pdf_page(tmp_path, "shx.pdf", n_words=20, n_lines=2000)
    ext = PlanBAVisionExtractor()
    with fitz.open(str(p_scan)) as d:
        assert ext.needs_vision_fallback(d[0]) is True
    with fitz.open(str(p_vec)) as d:
        assert ext.needs_vision_fallback(d[0]) is False
    with fitz.open(str(p_shx)) as d:
        assert ext.needs_vision_fallback(d[0]) is True


def test_render_page_high_res_cap(tmp_path):
    p = _pdf_page(tmp_path, "r.pdf", 20, 10)
    ext = PlanBAVisionExtractor(dpi=300, max_side_px=1500)
    with fitz.open(str(p)) as d:
        img = ext.render_page_high_res(d[0])
    assert img.ndim == 3 and img.shape[2] == 3
    assert max(img.shape[:2]) <= 1500


# ============================================================================
# Couche 2 : tableaux, contours, bulles d'axes
# ============================================================================

def _blank(w=1600, h=1200):
    return np.full((h, w, 3), 255, dtype=np.uint8)


def test_detect_table_rois_morphologie():
    img = _blank()
    x0, y0, cw, ch, nc, nr = 200, 150, 150, 60, 6, 8
    for c in range(nc + 1):
        cv2.line(img, (x0 + c * cw, y0), (x0 + c * cw, y0 + nr * ch), (0, 0, 0), 2)
    for r in range(nr + 1):
        cv2.line(img, (x0, y0 + r * ch), (x0 + nc * cw, y0 + r * ch), (0, 0, 0), 2)
    rois = PlanBAVisionExtractor().detect_table_and_schedule_rois(img)
    assert len(rois) >= 1
    x, y, w, h = rois[0]
    assert abs(x - x0) < 15 and abs(y - y0) < 15
    assert abs(w - nc * cw) < 30 and abs(h - nr * ch) < 30
    # Image vide : aucun tableau
    assert PlanBAVisionExtractor().detect_table_and_schedule_rois(_blank()) == []


def test_detect_structural_contours():
    img = _blank()
    boxes = [(300, 300, 80, 120), (700, 400, 100, 100), (1100, 600, 60, 150)]
    for x, y, w, h in boxes:
        cv2.rectangle(img, (x, y), (x + w, y + h), (0, 0, 0), 3)
    found = PlanBAVisionExtractor().detect_structural_elements_contours(img)
    for x, y, w, h in boxes:
        assert any(abs(fx - x) < 10 and abs(fy - y) < 10 and abs(fw - w) < 12
                   and abs(fh - h) < 12 for fx, fy, fw, fh in found), (x, y, w, h)


def test_grid_axes_bulles_et_affectation():
    img = _blank(1600, 1200)
    letters = {"A": (300, 80), "B": (700, 80), "C": (1100, 80)}
    numbers = {"1": (80, 300), "2": (80, 700)}
    ocr = []
    for t, (cx, cy) in {**letters, **numbers}.items():
        cv2.circle(img, (cx, cy), 36, (0, 0, 0), 3)
        ocr.append({"text": t, "bbox": (cx - 10, cy - 12, 20, 24), "conf": 0.95})
    ext = PlanBAVisionExtractor()
    ext.last_scale = 300 / 72.0
    bubbles = ext.detect_grid_axis_bubbles(img, ocr)
    found = {b["axis"] for b in bubbles}
    assert {"A", "B", "C", "1", "2"} <= found
    axes = ext.assign_grid_axes([(700, 300), (1100, 700), (300, 690)], bubbles)
    assert axes == [["B-1"], ["C-2"], ["A-2"]]


# ============================================================================
# Couche 3 : graphe spatial cKDTree + anti-hallucination
# ============================================================================

def test_association_spatiale_et_correction_decimale():
    ext = PlanBAVisionExtractor()
    ocr = [
        {"text": "P12", "bbox": (100, 100, 40, 20), "conf": 0.95},
        {"text": "25x400", "bbox": (100, 140, 60, 20), "conf": 0.90},   # erreur OCR
        {"text": "6HA14", "bbox": (100, 180, 60, 20), "conf": 0.93},
        {"text": "Cad HA8 e=15", "bbox": (100, 210, 110, 20), "conf": 0.91},
        {"text": "S3", "bbox": (1200, 900, 30, 20), "conf": 0.96},
        {"text": "140x140x40", "bbox": (1200, 930, 110, 20), "conf": 0.92},
        {"text": "LG7", "bbox": (3000, 3000, 40, 20), "conf": 0.9},     # trop loin de tout
    ]
    contours = [(60, 240, 80, 128)]   # poteau 25x40 -> ratio 0.625
    els = ext.associate_ocr_tokens_spatially(ocr, page_num=2, contours=contours)
    by = {e.label: e for e in els}
    assert set(by) == {"P12", "S3", "LG7"}

    p12 = by["P12"]
    assert p12.category == "poteau" and p12.page_num == 2
    assert p12.dimensions == {"b": 0.25, "h": 0.40}
    assert any("Correction décimale" in w for w in p12.warnings)
    assert "6HA14" in p12.rebar_spec and "Cad HA8 e=15" in p12.rebar_spec
    assert p12.source_engine == "cv_spatial_graph"

    s3 = by["S3"]
    assert s3.category == "semelle"
    assert s3.dimensions == {"A": 1.40, "B": 1.40, "H": 0.40}
    assert s3.rebar_spec == []

    lg7 = by["LG7"]
    assert lg7.category == "longrine" and lg7.dimensions == {} and lg7.rebar_spec == []


def test_sanity_check_sans_contour_conserve_valeurs_plausibles():
    dims, warn = PlanBAVisionExtractor.geometric_sanity_check({"b": 0.25, "h": 0.45}, None)
    assert dims == {"b": 0.25, "h": 0.45} and warn == []
    dims, warn = PlanBAVisionExtractor.geometric_sanity_check({"b": 0.25, "h": 0.45}, (50.0, 90.0))
    assert dims == {"b": 0.25, "h": 0.45} and warn == []


# ============================================================================
# Couche 3 bis : VLM (mock HTTP) + couche 4 : overlay
# ============================================================================

def test_vlm_router_mock_et_parse(monkeypatch):
    monkeypatch.setenv("PLANBA_VLM_BACKEND", "ollama")
    import requests

    class _Resp:
        status_code = 200
        text = ""

        def json(self):
            return {"choices": [{"message": {"content":
                '{"elements": [{"category": "poteau", "label": "P4", "quantity": 3, '
                '"dimensions": {"b": 0.25, "h": 0.40, "L": 3.0}, '
                '"rebar_spec": ["6HA14", "Cad HA8 e=15"], "grid_axes": ["A-1","B-1","C-1"]}, '
                '{"label": "", "category": "poutre"}, '
                '{"label": "N7", "category": "???", "dimensions": {"b": "0,20", "h": "abc"}}]}'
            }}]}

    captured = {}

    def fake_post(url, json=None, headers=None, timeout=None):
        captured["url"] = url
        captured["payload"] = json
        return _Resp()

    monkeypatch.setattr(requests, "post", fake_post)
    router = VLMRouter()
    assert router.available and router.backend.name == "ollama"
    ext = PlanBAVisionExtractor(vlm_router=router)
    img = _blank(800, 600)
    els = ext.extract_roi_with_vlm(img, (10, 10, 400, 300), page_num=3)

    assert captured["url"].endswith("/chat/completions")
    assert captured["payload"]["temperature"] == 0.0
    assert captured["payload"]["response_format"] == {"type": "json_object"}
    assert captured["payload"]["messages"][1]["content"][1]["image_url"]["url"].startswith(
        "data:image/png;base64,")

    assert [e.label for e in els] == ["P4", "N7"]
    p4 = els[0]
    assert p4.quantity == 3 and p4.dimensions == {"b": 0.25, "h": 0.40, "L": 3.0}
    assert p4.grid_axes == ["A-1", "B-1", "C-1"] and p4.source_engine == "cv_vlm_roi"
    assert p4.crop_image_bgr is not None and p4.bbox == (10, 10, 400, 300)
    n7 = els[1]
    assert n7.category == "poutre" and n7.dimensions == {"b": 0.20}


def test_vlm_router_off_degrade_proprement(monkeypatch):
    monkeypatch.setenv("PLANBA_VLM_BACKEND", "off")
    router = VLMRouter()
    assert not router.available
    ext = PlanBAVisionExtractor(vlm_router=router)
    assert ext.extract_roi_with_vlm(_blank(300, 300), (0, 0, 100, 100), 1) == []


def test_render_audit_overlay(tmp_path):
    img = _blank(900, 700)
    els = [
        DrawingElement("poteau", "P1", {"b": 0.25, "h": 0.4}, bbox=(50, 50, 80, 120), confidence=0.95),
        DrawingElement("semelle", "S1", {"A": 1.4}, bbox=(300, 300, 150, 150), confidence=0.70),
        DrawingElement("poutre", "N1", {}, bbox=(600, 100, 200, 40), confidence=0.99),
    ]
    out = PlanBAVisionExtractor().render_audit_overlay(
        img, els, anomaly_labels=["N1"], out_path=str(tmp_path / "ov" / "page_1.png"),
        table_rois=[(500, 400, 300, 200)])
    assert Path(out).exists()
    canvas = cv2.imread(out)
    assert canvas is not None and canvas.shape[:2] == (700, 900)
    assert tuple(int(v) for v in canvas[50, 90]) == (0, 180, 0)      # vert
    assert tuple(int(v) for v in canvas[300, 375]) == (0, 140, 255)  # orange
    assert tuple(int(v) for v in canvas[100, 700]) == (0, 0, 255)    # rouge


# ============================================================================
# Pont vers le pipeline legacy (plan_data) + contrats
# ============================================================================

def test_drawing_element_vers_record_et_schema():
    el = DrawingElement("poteau", "p3", {"b": 25, "h": 40}, ["6HA14"],
                        grid_axes=["A-1", "B-1"], page_num=2, bbox=(1, 2, 3, 4),
                        confidence=0.9, source_engine="cv_spatial_graph")
    assert el.label == "P3"
    rec = el.to_plan_record()
    assert rec.qty_plan_coffrage == 2 and rec.qty_table_schedule is None
    el_t = DrawingElement("poteau", "P3", {"b": 0.25}, quantity=4, source_engine="cv_vlm_roi")
    rec_t = el_t.to_plan_record()
    assert rec_t.qty_table_schedule == 4 and rec_t.qty_plan_coffrage is None
    sch = el.to_schema_element()
    assert sch.page == 2 and sch.bbox == (1.0, 2.0, 4.0, 6.0)
    assert sch.provenance["category"] == "poteau"


def test_pont_legacy_plan_data():
    from core.local_extractor import verifier_livrables_ou_lever
    pop = PlanBAExcelPopulator("Legacy")
    recs = pop.gatekeeper.deduplicate_records(_sample_records())
    plan = records_to_legacy_plan_data(
        recs, projet="Legacy", normalizer=pop.gatekeeper.normalize_dimensions_to_meters)
    assert verifier_livrables_ou_lever(plan) is True
    cat, impl = plan["catalogue_types"], plan["implantations"]
    assert cat["semelles"]["S1"]["a"] == 1.40 and cat["semelles"]["S1"]["h"] == 0.35
    assert cat["semelles"]["S1"]["ferr_x"] == {"nb": 12, "phi": 14}
    assert len(impl["semelles"]) == 4 and impl["semelles"][0]["axe"] == "A"
    assert cat["poteaux"]["P1"]["a"] == 0.30 and cat["poteaux"]["P1"]["cadres"] == {"phi": 8, "esp": 0.12}
    assert cat["longrines"]["LG1"]["portee"] == 4.80
    assert len(impl["poteaux"]) == 6


def test_porte2_deduplication_fusionne_sources():
    gate = PlanBAGatekeeper(MoroccanBETEngine())
    a = ExtractedPlanRecord("poteau", "P1", grid_axes=["A-1", "B-1"], qty_plan_coffrage=2,
                            raw_dimensions={"b": 25}, rebar_specs=["6HA14"])
    b = ExtractedPlanRecord("poteau", "p1", qty_table_schedule=2,
                            raw_dimensions={"b": 25, "h": 40}, rebar_specs=["6HA14", "Cad HA8 e=15"])
    merged = gate.deduplicate_records([a, b])
    assert len(merged) == 1
    m = merged[0]
    assert m.grid_axes == ["A-1", "B-1"]
    assert m.qty_plan_coffrage == 2 and m.qty_table_schedule == 2
    assert m.raw_dimensions == {"b": 25, "h": 40}
    assert m.rebar_specs == ["6HA14", "Cad HA8 e=15"]


def test_porte2_cumul_occurrences_coffrage():
    gate = PlanBAGatekeeper(MoroccanBETEngine())
    occ = [DrawingElement("poteau", "P1", {"b": 25, "h": 35}, page_num=1,
                          bbox=(i * 500, 100, 40, 20), source_engine="cv_spatial_graph")
           .to_plan_record() for i in range(8)]
    table = DrawingElement("poteau", "P1", {"b": 25, "h": 35}, ["6HA14"], quantity=8,
                           source_engine="cv_vlm_roi").to_plan_record()
    merged = gate.deduplicate_records(occ + [table])
    assert len(merged) == 1
    assert merged[0].qty_plan_coffrage == 8 and merged[0].qty_table_schedule == 8
    assert gate.effective_quantity(merged[0]) == 8
    one = ExtractedPlanRecord("poteau", "P9", grid_axes=["A-1"], qty_plan_coffrage=3)
    assert gate.effective_quantity(one) == 3


def test_etiquettes_couplees_et_triplet_semelle():
    ext = PlanBAVisionExtractor()
    ocr = [
        {"text": "(S1,P1)", "bbox": (100, 100, 70, 20), "conf": 0.95},
        {"text": "150x150x40", "bbox": (100, 130, 100, 20), "conf": 0.95},
    ]
    els = ext.associate_ocr_tokens_spatially(ocr, page_num=1)
    by = {e.label: e for e in els}
    assert set(by) == {"S1", "P1"}
    assert by["S1"].dimensions == {"A": 1.5, "B": 1.5, "H": 0.4}
    assert by["P1"].dimensions == {}          # le triplet appartient a la semelle
    assert PlanBAVisionExtractor._section_dims(("25", "45", "480"), "poutre") == {
        "b": 0.25, "h": 0.45, "L": 4.80}


def test_bornes_ratios_et_catalogue_conformes_spec():
    assert MOROCCAN_STEEL_RATIO_BOUNDS["semelle"] == (40.0, 95.0)
    assert MOROCCAN_STEEL_RATIO_BOUNDS["poteau"] == (100.0, 210.0)
    assert MOROCCAN_REBAR_CATALOG[32] == 6.313 and MOROCCAN_REBAR_CATALOG[40] == 9.864
    assert [c for c, *_ in BPU_ITEMS][:2] == ["B.01", "B.02"] and BPU_ITEMS[-1][0] == "A.01"
