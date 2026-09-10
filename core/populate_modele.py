#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
core/populate_modele.py -- Pont vers le gabarit standardise 5 feuilles.

Remplit le classeur modele (01_Detail_Quantitatif, 02_Armatures,
03_Attachement_Ferraillage, 04_GO_Attachement,
05_Catalogue_Armatures_Standard) a partir des donnees extraites.

Contrat d'entree (format plat) :
    {
      "projet": "Nom du projet",          # str
      "semelles": [                        # 1 entree par semelle IMPLANTEE
        {"type": "S1", "axe": "A", "file": "1",
         "a": 1.2, "b": 1.2, "h": 0.3,
         "phi": 12, "nb_x": 8, "nb_y": 8}, ...
      ],
      "poteaux": [                         # optionnel, 1 par instance
        {"type": "P1", "axe": "", "file": "",
         "a": 0.25, "b": 0.35, "hauteur": 3.0,
         "long_bars": [{"nb": 6, "phi": 14}],
         "cadres": {"phi": 6, "esp": 0.15}}, ...
      ],
      "poutres": [                         # optionnel, 1 par instance
        {"type": "N1", "axe": "", "b": 0.20, "h": 0.30,
         "portee": 4.0,                    # None si non cotee
         "filants_inf": [{"nb": 3, "phi": 14}],
         "filants_sup": [{"nb": 2, "phi": 12}],
         "cadres": {"phi": 6, "esp": 0.18}}, ...
      ],
    }

Allocation 100% dynamique : les blocs sont reconstruits selon le nombre
reel d'elements (aucun chevauchement, sous-totaux et liens recalcules).
Les styles sont recopies depuis le gabarit (aucune charte en dur ici).
"""
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

DIAMETRES_DEFAUT = [6, 8, 10, 12, 14, 16, 20, 25, 32]
POIDS_ML = [0.222, 0.395, 0.617, 0.888, 1.208, 1.578, 2.466, 3.853, 6.313]
ENROBAGE = 0.05
COEF_ANCRAGE = 34


def _safe_float(value, default=0.0):
    if value is None or isinstance(value, bool):
        return default
    try:
        return float(str(value).strip().replace(",", "."))
    except (TypeError, ValueError):
        return default


def _safe_int(value, default=0):
    try:
        return int(_safe_float(value, default))
    except (TypeError, ValueError, OverflowError):
        return default


def _safe_dict(value):
    return value if isinstance(value, dict) else {}


def _safe_list(value):
    return value if isinstance(value, list) else []


def _normalise_payload(plan_data):
    source = _safe_dict(plan_data)
    result = {"projet": source.get("projet", "Projet BTP")}
    for family in ("semelles", "poteaux", "poutres", "longrines",
                   "chainages", "murs", "voiles"):
        result[family] = []
        for raw in _safe_list(source.get(family)):
            item = _safe_dict(raw).copy()
            for name in ("a", "b", "h", "hauteur", "portee", "ep"):
                if name in item and item[name] is not None:
                    item[name] = _safe_float(item[name])
            for name in ("phi", "nb_x", "nb_y"):
                if name in item:
                    item[name] = _safe_int(item[name])
            for name in ("long_bars", "filants_inf", "filants_sup"):
                if name in item:
                    item[name] = [
                        {"nb": _safe_int(_safe_dict(bar).get("nb")),
                         "phi": _safe_int(_safe_dict(bar).get("phi"))}
                        for bar in _safe_list(item[name])
                    ]
            cadres = item.get("cadres")
            if isinstance(cadres, dict):
                cadres = cadres.copy()
                cadres["phi"] = _safe_int(cadres.get("phi"))
                cadres["esp"] = _safe_float(cadres.get("esp"), 0.15)
                item["cadres"] = cadres
            result[family].append(item)
    return result


# ============================================================================
# Utilitaires
# ============================================================================

def set_cell_safe(ws, row: int, col: int, value):
    """Ecriture tolerante aux cellules fusionnees (cartouche)."""
    cell = ws.cell(row, col)
    if type(cell).__name__ == "MergedCell":
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= row <= rng.max_row \
                    and rng.min_col <= col <= rng.max_col:
                ws.cell(rng.min_row, rng.min_col).value = value
                return
    else:
        cell.value = value


def _snapshot_row(ws, row, ncols):
    """Capture les styles d'une ligne modele pour les rejouer."""
    snap = []
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        snap.append({
            "font": copy(cell.font),
            "fill": copy(cell.fill),
            "border": copy(cell.border),
            "alignment": copy(cell.alignment),
            "number_format": cell.number_format,
        })
    return snap


def _apply_row(ws, row, snap, values):
    """Ecrit une ligne complete en appliquant un style capture."""
    for c, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=c)
        st = snap[c - 1] if c - 1 < len(snap) else snap[-1]
        cell.value = val
        cell.font = copy(st["font"])
        cell.fill = copy(st["fill"])
        cell.border = copy(st["border"])
        cell.alignment = copy(st["alignment"])
        cell.number_format = st["number_format"]


def _lire_colonnes_diam(ws, header_row=3, first_col=11):
    """Detecte les colonnes de ventilation (T6..T32) depuis les en-tetes."""
    cols = []
    for c in range(first_col, ws.max_column + 2):
        txt = str(ws.cell(row=header_row, column=c).value or "")
        txt = txt.replace(" ", "").upper()
        if txt.startswith("T") and txt[1:].isdigit():
            cols.append((c, int(txt[1:])))
        elif cols:
            break
    if not cols:
        cols = [(first_col + i, d) for i, d in enumerate(DIAMETRES_DEFAUT)]
    return cols


def _trouver_ligne_libelle(ws, libelle, col=1, debut=1):
    """Retrouve une ligne de recap par son libelle (insensible casse)."""
    cible = libelle.strip().lower()
    for r in range(debut, ws.max_row + 1):
        if str(ws.cell(row=r, column=col).value or "").strip().lower() \
                == cible:
            return r
    return None


def _clear_rows(ws, first, last):
    """Supprime un bloc de lignes (aucune fusion sous la ligne 5)."""
    if last >= first:
        ws.delete_rows(first, last - first + 1)


def _formules_acier(col_diam, row, col_g="G", col_h="H", col_j="J",
                    col_i="I"):
    """Ventilation =IF($I=diam, $G*$H*$J, "") pour chaque diametre."""
    return {c: f'=IF(${col_i}{row}={d},'
                f'${col_g}{row}*${col_h}{row}*${col_j}{row},"")'
            for c, d in col_diam}


# ============================================================================
# FEUILLE 1 : 01_Detail_Quantitatif
# ============================================================================

def _remplir_detail(ws, semelles, poteaux, poutres=None, longrines=None,
                    chainages=None, murs=None, voiles=None):
    """Reconstruit la zone de donnees (des la ligne 7).
    Retourne les lignes de sous-totaux pour les liens de la feuille 4."""
    NCOLS = 11
    st_entry = _snapshot_row(ws, 9, NCOLS)
    st_sub = _snapshot_row(ws, 17, NCOLS)
    st_section = _snapshot_row(ws, 7, NCOLS)
    st_total = _snapshot_row(ws, 42, NCOLS)
    _clear_rows(ws, 7, max(ws.max_row, 42))

    articles = [
        ("Terrassement en fouilles",
         [(s.get("axe", ""), s.get("file", ""),
           f"Fouille Semelle {s.get('type', '?')}", "M3", 1,
           round(float(s.get("a", 0)) + 0.20, 3),
           round(float(s.get("b", 0)) + 0.20, 3), 1.20)
          for s in semelles]),
        ("Béton de propreté dosé à 150 kg/m³",
         [(s.get("axe", ""), s.get("file", ""),
           f"BP Semelle {s.get('type', '?')}", "M3", 1,
           round(float(s.get("a", 0)) + 0.10, 3),
           round(float(s.get("b", 0)) + 0.10, 3), 0.10)
          for s in semelles]),
        ("Béton armé en fondation dosé à 350 kg/m³",
         [(s.get("axe", ""), s.get("file", ""),
           f"Semelle {s.get('type', '?')}", "M3", 1,
           float(s.get("a", 0)), float(s.get("b", 0)),
           float(s.get("h", 0)))
          for s in semelles]),
    ]
    if poteaux:
        articles.append(
            ("Béton pour fûts de poteaux en fondation",
             [(p.get("axe", ""), p.get("file", ""),
               f"Fût {p.get('type', '?')}", "M3", 1,
               float(p.get("a", 0)), float(p.get("b", 0)),
               float(p.get("hauteur", 3.0)))
              for p in poteaux]))

    # Poutres
    if poutres:
        articles.append(
            ("Béton pour poutres",
             [(p.get("axe", ""), "",
               f"Poutre {p.get('type', '?')}", "M3", 1,
               float(p.get("b", 0)), float(p.get("h", 0)),
               float(p.get("portee") or 3.5))
              for p in poutres]))

    # Longrines
    if longrines:
        articles.append(
            ("Béton pour longrines",
             [(l.get("axe", ""), "",
               f"Longrine {l.get('type', '?')}", "M3", 1,
               float(l.get("b", 0)), float(l.get("h", 0)),
               float(l.get("portee") or 3.5))
              for l in longrines]))

    # Chainages
    if chainages:
        articles.append(
            ("Béton pour chaînages",
             [(c.get("axe", ""), "",
               f"Chaînage {c.get('type', '?')}", "M3", 1,
               float(c.get("b", 0)), float(c.get("h", 0)),
               float(c.get("portee") or 3.5))
              for c in chainages]))

    # Murs (bandes noyees)
    if murs:
        articles.append(
            ("Béton pour bandes noyées",
             [(m.get("axe", ""), "",
               f"BN {m.get('type', '?')}", "M3", 1,
               float(m.get("b", 0)), float(m.get("h", 0)),
               float(m.get("portee") or 3.5))
              for m in murs]))

    # Voiles
    if voiles:
        articles.append(
            ("Béton pour voiles",
             [(v.get("axe", ""), "",
               f"Voile {v.get('type', '?')}", "M3", 1,
               float(v.get("ep", 0.20)),
               float(v.get("hauteur", 3.0)),
               1.0)
              for v in voiles]))

    row = 7
    subs = {}
    for num, (titre, lignes) in enumerate(articles, 1):
        vals = [num, None, None, titre] + [None] * 7
        _apply_row(ws, row, st_section, vals)
        row += 1
        vals = [None, "Axe", "File"] + [None] * 8
        _apply_row(ws, row, st_section, vals)
        row += 1
        debut = row
        for (axe, file_, des, unit, n, g, h, i) in lignes:
            j = (f"=IF(COUNTA(G{row}:I{row})>0,"
                 f"F{row}*PRODUCT(G{row}:I{row}),F{row})")
            vals = [None, axe, file_, des, unit, n, g, h, i, j, None]
            _apply_row(ws, row, st_entry, vals)
            row += 1
        fin = row - 1
        if lignes:
            sub_j = f"=SUM(J{debut}:J{fin})"
            sub_k = f"=SUM(J{debut}:J{fin})"
        else:
            sub_j = sub_k = 0
        vals = ["Ss-total"] + [None] * 8 + [sub_j, sub_k]
        _apply_row(ws, row, st_sub, vals)
        subs[num] = row
        row += 1

    tot_j = "+".join(f"J{r}" for r in subs.values())
    tot_k = "+".join(f"K{r}" for r in subs.values())
    vals = ["TOTAL GÉNÉRAL"] + [None] * 8 + [f"={tot_j}", f"={tot_k}"]
    _apply_row(ws, row, st_total, vals)
    return {"subs": subs, "total": row}


# ============================================================================
# FEUILLE 2 : 02_Armatures
# ============================================================================

def _remplir_armatures(ws, semelles, poteaux, poutres):
    """Reconstruit les lignes d'armatures (des la ligne 4) puis le recap.
    Retourne la cellule du poids total acier (ligne, colonne)."""
    NM = 10
    col_diam = _lire_colonnes_diam(ws)
    st_entry = _snapshot_row(ws, 4, NM + len(col_diam))
    # Styles du recap existant (recherche par libelle, repli neutre)
    st_recap = _snapshot_row(
        ws, _trouver_ligne_libelle(ws, "LONGUEUR TOTALE (ml)") or 21,
        NM + len(col_diam))
    st_pml = _snapshot_row(
        ws, _trouver_ligne_libelle(ws, "POIDS UNITAIRE NOMINAL (kg/ml)")
        or 22, NM + len(col_diam))
    st_total = _snapshot_row(
        ws, _trouver_ligne_libelle(ws, "POIDS TOTAL DES ACIERS (KG)")
        or 24, NM + len(col_diam))

    _clear_rows(ws, 4, ws.max_row)
    row = 4

    def _ligne_mere(vals_main):
        """Ligne element + ventilation (G/H/J vides -> "")."""
        nonlocal row
        vals = list(vals_main) + [None] * (
            NM + len(col_diam) - len(vals_main))
        _apply_row(ws, row, st_entry, vals)
        for c, d in col_diam:
            ws.cell(row=row, column=c,
                    value=f'=IF($I{row}={d},$G{row}*$H{row}*$J{row},"")')
        row += 1
        return row - 1

    def _ligne_barre(label, g, h, i, j_formula):
        nonlocal row
        vals = [label] + [None] * 5 + [g, h, i, j_formula]
        vals += [None] * (NM + len(col_diam) - len(vals))
        _apply_row(ws, row, st_entry, vals)
        for c, d in col_diam:
            ws.cell(row=row, column=c,
                    value=f'=IF($I{row}={d},$G{row}*$H{row}*$J{row},"")')
        row += 1

    # --- Semelles ---
    for s in semelles:
        lab = s.get("type", "?")
        a, b, h = float(s.get("a", 0)), float(s.get("b", 0)), \
            float(s.get("h", 0))
        phi = int(s.get("phi", 0))
        nb_x, nb_y = int(s.get("nb_x", 0)), int(s.get("nb_y", 0))
        r0 = _ligne_mere([f"{lab}", s.get("axe", ""), s.get("file", ""),
                          a, b, h])
        _ligne_barre("Armature INF X", 1, nb_x, phi,
                     f"=(E{r0}-{ENROBAGE})+{COEF_ANCRAGE}*I{row}/1000")
        r1 = row - 1
        _ligne_barre("Armature INF Y", 1, nb_y, phi,
                     f"=(D{r0}-{ENROBAGE})+{COEF_ANCRAGE}*I{row}/1000")
        _ = r1

    # --- Poteaux ---
    for p in poteaux:
        lab = p.get("type", "?")
        a, b = float(p.get("a", 0)), float(p.get("b", 0))
        haut = p.get("hauteur", 3.0)
        r0 = _ligne_mere([f"{lab}", p.get("axe", ""), p.get("file", ""),
                          a, b, haut])
        for lb in p.get("long_bars", []) or []:
            if lb.get("nb", 0) > 0 and lb.get("phi", 0) > 0:
                _ligne_barre("ARM LONG", 1, lb["nb"], lb["phi"],
                             f"=(F{r0}+0.50)")
        cad = p.get("cadres", {}) or {}
        if cad.get("phi", 0) > 0 and (cad.get("esp") or 0) > 0:
            esp = cad["esp"]
            _ligne_barre("CADRE",
                         f"=ROUNDUP(F{r0}/{esp},0)+2",
                         f"=ROUNDUP(F{r0}/{esp},0)+2",
                         cad["phi"],
                         f"=(2*(D{r0}+E{r0})+0.10)")

    # --- Poutres ---
    for p in poutres:
        lab = p.get("type", "?")
        b, h = float(p.get("b", 0)), float(p.get("h", 0))
        portee = p.get("portee")
        r0 = _ligne_mere([f"{lab}", p.get("axe", ""), "",
                          b, h, portee if portee else None])
        for fi in (p.get("filants_inf", []) or []) + \
                (p.get("filants_sup", []) or []):
            if fi.get("nb", 0) > 0 and fi.get("phi", 0) > 0:
                j = (f"=(F{r0}+2*{COEF_ANCRAGE}*I{row}/1000)"
                     if portee else None)
                _ligne_barre("Armature", 1, fi["nb"], fi["phi"], j)
        cad = p.get("cadres", {}) or {}
        if cad.get("phi", 0) > 0 and (cad.get("esp") or 0) > 0 and portee:
            _ligne_barre("CADRE",
                         f"=ROUNDUP(F{r0}/{cad['esp']},0)+2",
                         f"=ROUNDUP(F{r0}/{cad['esp']},0)+2",
                         cad["phi"],
                         f"=(2*(D{r0}+E{r0})+0.10)")

    fin = row - 1

    # --- Recapitulatif ---
    def _recap(label, kind):
        nonlocal row
        vals = [label] + [None] * (NM + len(col_diam) - 1)
        _apply_row(ws, row,
                   st_total if kind in ("total",) else st_recap, vals)
        r = row
        row += 1
        return r

    col_k = col_diam[0][0]
    col_last = col_diam[-1][0]
    from openpyxl.utils import get_column_letter
    lettres = {d: get_column_letter(c) for c, d in col_diam}

    r_long = _recap("LONGUEUR TOTALE (ml)", "sum")
    for c, d in col_diam:
        ws.cell(row=r_long, column=c,
                value=f"=SUM({lettres[d]}4:{lettres[d]}{fin})")
    r_pml = _recap("POIDS UNITAIRE NOMINAL (kg/ml)", "pml")
    for (c, d), pml in zip(col_diam, POIDS_ML):
        ws.cell(row=r_pml, column=c, value=pml)
    r_part = _recap("POIDS PARTIEL PAR DIAMÈTRE (kg)", "part")
    for c, d in col_diam:
        ws.cell(row=r_part, column=c,
                value=f"={lettres[d]}{r_long}*{lettres[d]}{r_pml}")
    r_tot = _recap("POIDS TOTAL DES ACIERS (KG)", "total")
    ws.cell(row=r_tot, column=col_k,
            value=f"=SUM({lettres[col_diam[0][1]]}{r_part}:"
                  f"{lettres[col_diam[-1][1]]}{r_part})")
    r_ton = _recap("POIDS TOTAL (tonnes)", "total")
    ws.cell(row=r_ton, column=col_k,
            value=f"={lettres[col_diam[0][1]]}{r_tot}/1000")
    r_bet = _recap("Volume béton de référence (m³) — à saisir", "sum")
    r_taux = _recap("TAUX DE FERRAILLAGE MOYEN (kg/m³)", "total")
    ws.cell(row=r_taux, column=col_k,
            value=f"=IF({lettres[col_diam[0][1]]}{r_bet}>0,"
                  f"{lettres[col_diam[0][1]]}{r_tot}/"
                  f"{lettres[col_diam[0][1]]}{r_bet},\"\")")
    return {"total_kg_row": r_tot, "total_kg_col": col_k,
            "beton_row": r_bet, "data_fin": fin}


# ============================================================================
# FEUILLE 3 : 03_Attachement_Ferraillage
# ============================================================================

def _remplir_synthese(ws, semelles, poteaux, poutres):
    """Lignes de synthese par element + recap (longueurs, poids)."""
    NM = 6
    col_diam = _lire_colonnes_diam(ws, header_row=3, first_col=NM + 1)
    st_entry = _snapshot_row(ws, 4, NM + len(col_diam))
    _clear_rows(ws, 4, ws.max_row)
    row = 4

    def _ligne_syn(label, kind, n_ouv, nb, phi, long_u):
        nonlocal row
        vals = [label, kind, n_ouv, nb, phi, long_u]
        vals += [None] * (NM + len(col_diam) - len(vals))
        _apply_row(ws, row, st_entry, vals)
        for c, d in col_diam:
            ws.cell(row=row, column=c,
                    value=f'=IF($E{row}={d},$C{row}*$D{row}*$F{row},"")')
        row += 1

    for s in semelles:
        lab, phi = s.get("type", "?"), int(s.get("phi", 0))
        a = float(s.get("a", 0))
        long_u = round((a - ENROBAGE) + COEF_ANCRAGE * phi / 1000, 3) \
            if phi > 0 else 0
        nb = int(s.get("nb_x", 0)) + int(s.get("nb_y", 0))
        _ligne_syn(f"Semelle {lab}", "Lit INF X & Y", 1, nb, phi, long_u)
    for p in poteaux:
        for lb in p.get("long_bars", []) or []:
            if lb.get("nb", 0) > 0 and lb.get("phi", 0) > 0:
                haut = p.get("hauteur", 3.0)
                _ligne_syn(f"Poteau {p.get('type', '?')}", "ARM LONG",
                           1, lb["nb"], lb["phi"], round(haut + 0.50, 3))
    for p in poutres:
        portee = p.get("portee")
        for fi in (p.get("filants_inf", []) or []) + \
                (p.get("filants_sup", []) or []):
            if fi.get("nb", 0) > 0 and fi.get("phi", 0) > 0:
                lu = round(portee + 0.50, 3) if portee else ""
                _ligne_syn(f"Poutre {p.get('type', '?')}", "Filants",
                           1, fi["nb"], fi["phi"], lu)
    fin = row - 1

    from openpyxl.utils import get_column_letter
    lettres = {d: get_column_letter(c) for c, d in col_diam}
    vals = ["LONGUEUR TOTALE (ml)"] + [None] * (NM + len(col_diam) - 1)
    _apply_row(ws, row, st_entry, vals)
    r_long = row
    for c, d in col_diam:
        ws.cell(row=row, column=c,
                value=f"=SUM({lettres[d]}4:{lettres[d]}{fin})")
    row += 1
    vals = ["POIDS TOTAL SYNTHÈSE (kg)"] + [None] * (NM + len(col_diam) - 1)
    _apply_row(ws, row, st_entry, vals)
    r_poids = row
    parts = "+".join(f"{lettres[d]}{r_long}*{pml}"
                     for (c, d), pml in zip(col_diam, POIDS_ML))
    ws.cell(row=row, column=NM + 1, value=f"={parts}")
    return {"poids_row": r_poids}


# ============================================================================
# FEUILLE 4 : 04_GO_Attachement
# ============================================================================

def _mettre_a_jour_bordereau(ws, liens):
    """Met a jour les liens dynamiques (col E) + ligne futs si presents.
    liens = {"terr": k, "prop": k, "ba": k, "futs": k|None,
             "total_acier": "K24"/cellule, "futs_row": bool}"""
    set_cell_safe(ws, 4, 5, f"='01_Detail_Quantitatif'!K{liens['terr']}")
    set_cell_safe(ws, 5, 5, f"='01_Detail_Quantitatif'!K{liens['prop']}")
    set_cell_safe(ws, 6, 5, f"='01_Detail_Quantitatif'!K{liens['ba']}")
    set_cell_safe(ws, 7, 5, f"='02_Armatures'!{liens['total_acier']}")
    if liens.get("futs") is not None:
        r = 8
        set_cell_safe(ws, r, 1, "5")
        set_cell_safe(ws, r, 2, "Béton pour fûts de poteaux en fondation")
        set_cell_safe(ws, r, 3, "M3")
        set_cell_safe(ws, r, 5,
                      f"='01_Detail_Quantitatif'!K{liens['futs']}")
        set_cell_safe(ws, r, 6, f"=D{r}-E{r}")
        set_cell_safe(ws, r, 7, f"=IF(D{r}>0,E{r}/D{r},0)")
        for c in range(1, 8):
            cell = ws.cell(row=r, column=c)
            if not cell.font or cell.font == openpyxl.styles.Font():
                pass


# ============================================================================
# Point d'entree
# ============================================================================

def injecter_metre_dans_modele(plan_data: dict, template_path: str,
                               output_path: str):
    """Remplit le gabarit 5 feuilles. Retourne le chemin de sortie."""
    for feuille in ("01_Detail_Quantitatif", "02_Armatures",
                    "03_Attachement_Ferraillage", "04_GO_Attachement"):
        pass  # controle d'existence ci-dessous
    wb = openpyxl.load_workbook(template_path)
    manquantes = [f for f in ("01_Detail_Quantitatif", "02_Armatures",
                              "03_Attachement_Ferraillage",
                              "04_GO_Attachement") if f not in wb.sheetnames]
    if manquantes:
        raise ValueError(
            f"Gabarit incomplet ({template_path}) : feuilles manquantes "
            f"{manquantes}")

    plan_data = _normalise_payload(plan_data)
    semelles = plan_data["semelles"]
    poteaux = plan_data["poteaux"]
    poutres = plan_data.get("poutres", [])
    longrines = plan_data.get("longrines", [])
    chainages = plan_data.get("chainages", [])
    murs = plan_data.get("murs", [])
    voiles = plan_data.get("voiles", [])
    nom_projet = plan_data.get("projet", "Projet BTP")

    ws1 = wb["01_Detail_Quantitatif"]
    for r in range(1, 6):
        for c in range(1, ws1.max_column + 1):
            if "projet" in str(ws1.cell(r, c).value or "").lower():
                set_cell_safe(ws1, r, c + 1, nom_projet)
                break

    info1 = _remplir_detail(ws1, semelles, poteaux, poutres, longrines,
                            chainages, murs, voiles)
    from openpyxl.utils import get_column_letter
    info2 = _remplir_armatures(ws1 and wb["02_Armatures"], semelles,
                               poteaux, poutres)
    _remplir_synthese(wb["03_Attachement_Ferraillage"], semelles,
                      poteaux, poutres)

    col_k_aciers = get_column_letter(info2["total_kg_col"])
    # Lien beton de reference -> total general feuille 1
    _mettre_a_jour_bordereau(
        wb["04_GO_Attachement"],
        {"terr": info1["subs"][1], "prop": info1["subs"][2],
         "ba": info1["subs"][3],
         "futs": info1["subs"].get(4),
         "total_acier": f"{col_k_aciers}{info2['total_kg_row']}"})
    ws2b = wb["02_Armatures"]
    set_cell_safe(ws2b, info2["beton_row"], info2["total_kg_col"],
                  f"='01_Detail_Quantitatif'!K{info1['total']}")

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    n_arm = len(semelles) * 2 + sum(
        len(p.get("long_bars", []) or []) for p in poteaux)
    print(f"Metre 5 feuilles genere : {output_path} "
          f"({len(semelles)} semelles, {len(poteaux)} poteaux, "
          f"{len(poutres)} poutres, {n_arm} lignes d'armatures)")
    return output_path


if __name__ == "__main__":
    sample_data = {
        "projet": "Projet R+2 Maroc",
        "semelles": [
            {"type": "S1", "axe": "A", "file": "1", "a": 1.2, "b": 1.2,
             "h": 0.3, "phi": 12, "nb_x": 8, "nb_y": 8},
            {"type": "S2", "axe": "B", "file": "2", "a": 1.5, "b": 1.5,
             "h": 0.4, "phi": 12, "nb_x": 10, "nb_y": 10}
        ]
    }
    injecter_metre_dans_modele(sample_data, "output/modele_metre_BA.xlsx",
                               "output/test_metre_complet.xlsx")
