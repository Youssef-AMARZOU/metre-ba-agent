#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
core/surface_metre_excel.py -- Classeur "Metre des surfaces" (plans archi).

Feuilles (formules Excel dynamiques, aucune valeur figee) :
  1_Surfaces_Pieces  : par etage et par piece, L x l, surface explicite
                       (=ROUND(D*E,2)), surface segmentee, surface retenue
                       (=IF(F>0,F,G)), perimetre, confiance, observations ;
                       sous-totaux par etage (SUMIF) et total general.
  2_Second_Oeuvre    : lots marocains derives (revetement sol, plinthes,
                       enduit + peinture murs = perimetre x HSP, faux plafond)
                       avec PU HT (MAD) modifiables, HT / TVA / TTC.
  3_Audit_Echelle    : estimations d'echelle, alertes, source OCR.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Sequence, Union

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SURFACE_SHEETS = ("1_Surfaces_Pieces", "2_Second_Oeuvre", "3_Audit_Echelle")

LOTS_SECOND_OEUVRE = [
    # code, designation, unite, base ('surface' | 'perimetre_hsp' | 'perimetre'), PU HT MAD
    ("SO.01", "Chape + revêtement de sol (carrelage / grès) y compris lit de pose", "m2", "surface", 220.0),
    ("SO.02", "Plinthes assorties (hauteur 7 à 10 cm)", "ml", "perimetre", 35.0),
    ("SO.03", "Enduit intérieur au mortier de ciment sur murs (HSP)", "m2", "perimetre_hsp", 60.0),
    ("SO.04", "Peinture vinylique intérieure sur murs (2 couches)", "m2", "perimetre_hsp", 45.0),
    ("SO.05", "Faux plafond / enduit et peinture sous plafond", "m2", "surface", 80.0),
]


class SurfaceMetreExcel:
    def __init__(self, project_name: str, hauteur_sous_plafond_m: float = 2.80,
                 tva_rate: float = 0.20):
        self.project_name = project_name
        self.hsp = hauteur_sous_plafond_m
        self.tva = tva_rate
        self.font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        self.font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        self.font_bold = Font(name="Calibri", size=10, bold=True)
        self.font_normal = Font(name="Calibri", size=10)
        self.font_formula = Font(name="Calibri", size=10, bold=True, color="0B5394")
        self.fill_navy = PatternFill("solid", fgColor="1B365D")
        self.fill_sub = PatternFill("solid", fgColor="2C5282")
        self.fill_zebra = PatternFill("solid", fgColor="F7FAFC")
        self.fill_total = PatternFill("solid", fgColor="E2E8F0")
        self.fill_ok = PatternFill("solid", fgColor="C6F6D5")
        self.fill_warn = PatternFill("solid", fgColor="FEEBC8")
        self.fill_err = PatternFill("solid", fgColor="FED7D7")
        thin = Side(border_style="thin", color="CBD5E0")
        self.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        self.layout: Dict[str, Any] = {}

    # ------------------------------------------------------------------
    def export(self, reports: Sequence[Any], output_path: Union[str, Path]) -> Path:
        wb = Workbook()
        wb.remove(wb.active)
        self._sheet_pieces(wb, reports)
        self._sheet_second_oeuvre(wb)
        self._sheet_audit(wb, reports)
        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(out))
        return out

    # ------------------------------------------------------------------
    def _sheet_pieces(self, wb, reports) -> None:
        ws = wb.create_sheet(SURFACE_SHEETS[0])
        self._banner(ws, f"MÉTRÉ DES SURFACES PAR PIÈCE — {self.project_name}",
                     "Surface retenue = cote explicite 'L x l' si présente, sinon segmentation calibrée", 11)
        headers = ["Étage / Niveau", "Pièce", "Source", "Longueur (m)", "Largeur (m)",
                   "Surface explicite (m²)", "Surface segmentée (m²)",
                   "Surface retenue (m²)", "Périmètre (m)", "Confiance", "Observations"]
        self._headers(ws, 5, headers)
        r = 6
        start = r
        levels: List[str] = []
        for rep in reports:
            levels.append(rep.level)
            for i, room in enumerate(rep.rooms):
                vals = [
                    rep.level, room.label, room.source,
                    room.length_m if room.length_m is not None else "",
                    room.width_m if room.width_m is not None else "",
                    f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})), ROUND(D{r}*E{r}, 2), "")',
                    room.area_segmented_m2 if room.area_segmented_m2 is not None else "",
                    f'=IF(ISNUMBER(F{r}), F{r}, IF(ISNUMBER(G{r}), G{r}, 0))',
                    room.perimeter_m if room.perimeter_m is not None else "",
                    room.confidence,
                    " | ".join(room.warnings) if room.warnings else "",
                ]
                self._row(ws, r, vals, i % 2 == 1, {6, 8})
                for c in (4, 5, 6, 7, 8, 9):
                    ws.cell(row=r, column=c).number_format = "0.00"
                ws.cell(row=r, column=10).number_format = "0.00"
                r += 1
        end = max(start, r - 1)
        ws.conditional_formatting.add(f"J{start}:J{end}", FormulaRule(
            formula=[f"$J{start}<0.5"], fill=self.fill_err))
        ws.conditional_formatting.add(f"J{start}:J{end}", FormulaRule(
            formula=[f"AND($J{start}>=0.5,$J{start}<0.85)"], fill=self.fill_warn))
        ws.conditional_formatting.add(f"J{start}:J{end}", FormulaRule(
            formula=[f"$J{start}>=0.85"], fill=self.fill_ok))

        # Sous-totaux par etage
        r = end + 2
        ws.cell(row=r, column=1, value="SOUS-TOTAUX PAR ÉTAGE").font = self.font_bold
        r += 1
        sub_start = r
        for lvl in dict.fromkeys(levels):
            ws.cell(row=r, column=1, value=lvl).font = self.font_bold
            c = ws.cell(row=r, column=8,
                        value=f"=ROUND(SUMIF($A${start}:$A${end}, A{r}, $H${start}:$H${end}), 2)")
            c.font = self.font_formula
            c.number_format = "0.00"
            p = ws.cell(row=r, column=9,
                        value=f"=ROUND(SUMIF($A${start}:$A${end}, A{r}, $I${start}:$I${end}), 2)")
            p.font = self.font_formula
            p.number_format = "0.00"
            for col in range(1, 12):
                ws.cell(row=r, column=col).fill = self.fill_total
            r += 1
        tot = r
        ws.cell(row=tot, column=1, value="TOTAL GÉNÉRAL SURFACES (m²)").font = self.font_bold
        t = ws.cell(row=tot, column=8, value=f"=ROUND(SUM(H{start}:H{end}), 2)")
        t.font = self.font_formula
        t.fill = self.fill_ok
        t.number_format = "0.00"
        tp = ws.cell(row=tot, column=9, value=f"=ROUND(SUM(I{start}:I{end}), 2)")
        tp.font = self.font_formula
        tp.number_format = "0.00"
        ws.freeze_panes = "A6"
        self._autofit(ws)
        self.layout["pieces"] = {"start": start, "end": end, "total_row": tot,
                                 "sub_start": sub_start}

    def _sheet_second_oeuvre(self, wb) -> None:
        ws = wb.create_sheet(SURFACE_SHEETS[1])
        self._banner(ws, f"SECOND ŒUVRE — LOTS DÉRIVÉS DES SURFACES — {self.project_name}",
                     "Quantités liées à la feuille 1 ; PU HT (MAD) à ajuster selon le CPS", 7)
        ws.cell(row=4, column=1, value="Hauteur sous plafond (m) :").font = self.font_bold
        hsp = ws.cell(row=4, column=2, value=self.hsp)
        hsp.number_format = "0.00"
        hsp.fill = self.fill_warn
        headers = ["N° Prix", "Désignation", "Unité", "Quantité",
                   "PU HT (MAD)", "Montant HT (MAD)", "Base de calcul"]
        self._headers(ws, 5, headers)
        lay = self.layout["pieces"]
        s1 = SURFACE_SHEETS[0]
        tot_surface = f"'{s1}'!$H${lay['total_row']}"
        tot_perim = f"'{s1}'!$I${lay['total_row']}"
        start = 6
        for i, (code, desc, unit, base, pu) in enumerate(LOTS_SECOND_OEUVRE):
            r = start + i
            if base == "surface":
                qty, note = f"=ROUND({tot_surface}, 2)", "Σ surfaces retenues"
            elif base == "perimetre":
                qty, note = f"=ROUND({tot_perim}, 2)", "Σ périmètres"
            else:
                qty, note = f"=ROUND({tot_perim}*$B$4, 2)", "Σ périmètres × HSP"
            self._row(ws, r, [code, desc, unit, qty, pu, f"=ROUND(D{r}*E{r}, 2)", note],
                      i % 2 == 1, {4, 6})
            ws.cell(row=r, column=4).number_format = "0.00"
            ws.cell(row=r, column=5).number_format = "#,##0.00"
            ws.cell(row=r, column=6).number_format = "#,##0.00"
        end = start + len(LOTS_SECOND_OEUVRE) - 1
        t_ht, t_tva, t_ttc = end + 1, end + 2, end + 3
        ws.cell(row=t_ht, column=2, value="TOTAL HT (MAD)").font = self.font_bold
        ws.cell(row=t_ht, column=6, value=f"=ROUND(SUM(F{start}:F{end}), 2)").font = self.font_formula
        ws.cell(row=t_tva, column=2, value=f"TVA ({self.tva*100:.0f} %)").font = self.font_bold
        ws.cell(row=t_tva, column=6, value=f"=ROUND(F{t_ht}*{self.tva}, 2)").font = self.font_formula
        ws.cell(row=t_ttc, column=2, value="TOTAL TTC (MAD)").font = self.font_bold
        ws.cell(row=t_ttc, column=6, value=f"=ROUND(F{t_ht}+F{t_tva}, 2)").font = self.font_formula
        for rr in (t_ht, t_tva, t_ttc):
            ws.cell(row=rr, column=6).number_format = "#,##0.00"
            for c in range(1, 8):
                ws.cell(row=rr, column=c).fill = self.fill_total
        self._autofit(ws)
        self.layout["so"] = {"start": start, "end": end, "t_ht": t_ht,
                             "t_tva": t_tva, "t_ttc": t_ttc}

    def _sheet_audit(self, wb, reports) -> None:
        ws = wb.create_sheet(SURFACE_SHEETS[2])
        self._banner(ws, f"AUDIT DE L'ÉCHELLE ET DES ALERTES — {self.project_name}",
                     "Calibration px/m par chaînes de cotes et cotes hors-tout", 6)
        self._headers(ws, 5, ["Page / Étage", "Méthode", "px / m", "Cote A", "Cote B", "Retenue"])
        r = 6
        for rep in reports:
            for e in rep.scale_estimates:
                self._row(ws, r, [f"p{rep.page_num} {rep.level}", e.get("method"),
                                  e.get("px_per_m"), e.get("a"), e.get("b"),
                                  rep.scale_px_per_m])
                r += 1
            ws.cell(row=r, column=1, value=f"Échelle retenue p{rep.page_num}").font = self.font_bold
            ws.cell(row=r, column=3, value=rep.scale_px_per_m)
            ws.cell(row=r, column=4, value=f"confiance {rep.scale_confidence}")
            ws.cell(row=r, column=5, value=f"OCR : {rep.ocr_engine or '-'} ({rep.n_ocr_tokens} tokens)")
            r += 2
        ws.cell(row=r, column=1, value="ALERTES").font = self.font_bold
        r += 1
        n_alert = 0
        for rep in reports:
            for w in rep.warnings:
                c = ws.cell(row=r, column=1, value=f"p{rep.page_num} {rep.level}")
                m = ws.cell(row=r, column=2, value=w)
                if w.startswith("[ALERTE"):
                    c.fill = self.fill_err
                    m.fill = self.fill_err
                r += 1
                n_alert += 1
        if n_alert == 0:
            ws.cell(row=r, column=1, value="Aucune alerte d'échelle.").fill = self.fill_ok
        self._autofit(ws)

    # ------------------------------------------------------------------
    def _banner(self, ws, title, subtitle, max_col):
        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=max_col)
        c = ws.cell(row=1, column=1, value=title)
        c.font, c.fill = self.font_title, self.fill_navy
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max_col)
        s = ws.cell(row=3, column=1, value=subtitle)
        s.font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
        s.fill = self.fill_sub
        s.alignment = Alignment(horizontal="center", vertical="center")

    def _headers(self, ws, row, headers):
        ws.row_dimensions[row].height = 30
        for col, h in enumerate(headers, start=1):
            c = ws.cell(row=row, column=col, value=h)
            c.font, c.fill, c.border = self.font_header, self.fill_navy, self.border
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _row(self, ws, row, values, zebra=False, formula_cols=None):
        formula_cols = formula_cols or set()
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = self.font_formula if col in formula_cols else self.font_normal
            c.border = self.border
            if zebra:
                c.fill = self.fill_zebra
            right = isinstance(v, (int, float)) or col in formula_cols
            c.alignment = Alignment(horizontal="right" if right else "left",
                                    vertical="center", wrap_text=(col == 11))

    @staticmethod
    def _autofit(ws):
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = 12
            for cell in col:
                if cell.row <= 3:
                    continue
                s = str(cell.value or "")
                width = max(width, 14 if s.startswith("=") else min(len(s) + 3, 60))
            ws.column_dimensions[letter].width = width
