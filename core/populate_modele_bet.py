#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
core/populate_modele_bet.py -- Classeur certifie BET Maroc (5 feuilles).

Toutes les quantites affichees sont des FORMULES Excel anglaises dynamiques
(ROUND / SUM / SUMIF / IF / AND) ; les valeurs Python calculees par
`MoroccanBETEngine` servent de temoin pour l'audit arithmetique.

Feuilles :
  1_Avant_Metre_BA           : J = ROUND(F*G*H*I, 3), total ROUND(SUM)
  2_Nomenclature_Acier       : H = ROUND(E*F*G, 2), ventilation HA6..HA32
  3_Bordereau_DQE            : SUMIF par N° prix, HT / TVA 20 % / TTC (MAD)
  4_Controle_BET_Ratios      : ratios kg/m3 vs RPS 2011 + anomalies + vignettes
  5_Optimisation_Barres_12m  : First-Fit Decreasing sur barres de 12.00 m
"""
from __future__ import annotations

import io
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.bet_contracts import (
    AuditAnomaly,
    AvantMetreLine,
    ExtractedPlanRecord,
    RebarScheduleRow,
)
from core.moroccan_bet_engine import (
    MOROCCAN_REBAR_CATALOG,
    MOROCCAN_STEEL_RATIO_BOUNDS,
    PRIX_CODES,
    MoroccanBETEngine,
    PlanBAGatekeeper,
)

__all__ = [
    "PlanBAExcelPopulator",
    "optimize_12m_stock_bars",
    "SHEET_NAMES",
    "BPU_ITEMS",
]

SHEET_NAMES = (
    "1_Avant_Metre_BA",
    "2_Nomenclature_Acier",
    "3_Bordereau_DQE",
    "4_Controle_BET_Ratios",
    "5_Optimisation_Barres_12m",
)

DATA_RANGE_END = 500  # plage dynamique des SUMIF inter-feuilles

# Bordereau des prix unitaires (MAD HT) — valeurs indicatives a ajuster au CPS
BPU_ITEMS: List[Tuple[str, str, str, float, str]] = [
    ("B.01", "Béton de propreté B10/B15 dosé à 250 kg/m³ sous ouvrages BA",
     "m3", 650.0, "DGA Art. 61"),
    ("B.02", "Béton armé B25/B30 en fondation pour semelles isolées et filantes",
     "m3", 1150.0, "NM 10.1.008"),
    ("B.04", "Béton armé B25 en fondation pour longrines et chaînages",
     "m3", 1200.0, "NM 10.1.008"),
    ("B.05", "Béton armé B25/B30 en élévation pour poteaux",
     "m3", 1300.0, "BAEL 91 / RPS 2011"),
    ("B.06", "Béton armé B25/B30 en élévation pour poutres",
     "m3", 1300.0, "BAEL 91 / RPS 2011"),
    ("B.07", "Béton armé B25/B30 pour voiles porteurs et murs de soutènement",
     "m3", 1350.0, "RPS 2000 v2011"),
    ("B.08", "Béton armé B25/B30 pour dalles pleines",
     "m3", 1300.0, "BAEL 91"),
    ("A.01", "Acier à haute adhérence FeE500 (fourniture, façonnage et pose)",
     "kg", 13.50, "NM 01.4.097"),
]


# ============================================================================
# Optimisation de coupe (FFD) sur barres commerciales de 12.00 m
# ============================================================================

def optimize_12m_stock_bars(schedule_rows: List[RebarScheduleRow],
                            stock_length_m: float = 12.00
                            ) -> Dict[int, Dict[str, Any]]:
    """First-Fit Decreasing par diametre : nb de barres de 12 m, chutes."""
    bars_by_dia: Dict[int, List[float]] = defaultdict(list)
    for row in schedule_rows:
        total_bars = row.nb_elements * row.nb_bars_per_elem
        if total_bars <= 0:
            continue
        rem = row.developed_length_m
        while rem > stock_length_m:
            bars_by_dia[row.diameter_mm].extend([stock_length_m] * total_bars)
            rem = round(rem - stock_length_m, 3)
        if rem > 0.05:
            bars_by_dia[row.diameter_mm].extend([rem] * total_bars)

    results: Dict[int, Dict[str, Any]] = {}
    for dia, lengths in sorted(bars_by_dia.items()):
        lengths.sort(reverse=True)
        bins: List[float] = []
        for cut in lengths:
            for i, free in enumerate(bins):
                if free + 1e-9 >= cut:
                    bins[i] = round(free - cut, 3)
                    break
            else:
                bins.append(round(stock_length_m - cut, 3))
        nb = len(bins)
        purchased = nb * stock_length_m
        useful = round(sum(lengths), 2)
        waste = round(purchased - useful, 2)
        results[dia] = {
            "nb_stock_bars_12m": nb,
            "useful_length_m": useful,
            "purchased_length_m": round(purchased, 2),
            "waste_length_m": waste,
            "waste_pct": round(waste / purchased * 100.0, 2) if purchased else 0.0,
            "purchased_weight_kg": round(
                purchased * MOROCCAN_REBAR_CATALOG.get(dia, 0.0), 2),
        }
    return results


# ============================================================================
# Generateur Excel
# ============================================================================

class PlanBAExcelPopulator:
    """Pipeline : Gatekeeper (4 portes) -> MoroccanBETEngine -> 5 feuilles."""

    DIAMETERS_COLS = [6, 8, 10, 12, 14, 16, 20, 25, 32]
    CONTROL_CATEGORIES = ["semelle", "longrine", "chainage", "poteau",
                          "poutre", "voile", "dalle_pleine"]

    def __init__(self, project_name: str,
                 bet_name: str = "PlanBA - Métré Extracteur",
                 engine: Optional[MoroccanBETEngine] = None,
                 tva_rate: float = 0.20):
        self.project_name = project_name
        self.bet_name = bet_name
        self.tva_rate = tva_rate
        self.engine = engine or MoroccanBETEngine()
        self.gatekeeper = PlanBAGatekeeper(self.engine)

        # Resultats du dernier `process_and_export` (temoins d'audit)
        self.avant_metre_lines: List[AvantMetreLine] = []
        self.rebar_schedule_rows: List[RebarScheduleRow] = []
        self.records: List[ExtractedPlanRecord] = []
        self.layout: Dict[str, Any] = {}

        self.font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        self.font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        self.font_bold = Font(name="Calibri", size=10, bold=True)
        self.font_normal = Font(name="Calibri", size=10)
        self.font_formula = Font(name="Calibri", size=10, bold=True, color="0B5394")
        self.fill_navy = PatternFill("solid", fgColor="1B365D")
        self.fill_subhdr = PatternFill("solid", fgColor="2C5282")
        self.fill_zebra = PatternFill("solid", fgColor="F7FAFC")
        self.fill_total = PatternFill("solid", fgColor="E2E8F0")
        self.fill_ok = PatternFill("solid", fgColor="C6F6D5")
        self.fill_warn = PatternFill("solid", fgColor="FEEBC8")
        self.fill_err = PatternFill("solid", fgColor="FED7D7")
        thin = Side(border_style="thin", color="CBD5E0")
        medium = Side(border_style="medium", color="1B365D")
        self.border_cell = Border(left=thin, right=thin, top=thin, bottom=thin)
        self.border_total = Border(top=thin, bottom=medium, left=thin, right=thin)

    # ------------------------------------------------------------------
    # Calcul (deterministe) a partir des enregistrements
    # ------------------------------------------------------------------
    def compute(self, raw_records: List[ExtractedPlanRecord]
                ) -> Tuple[List[AvantMetreLine], List[RebarScheduleRow]]:
        """Porte 2 -> Porte 1 -> Porte 3 -> moteur -> Porte 4."""
        self.engine.reset()
        records = self.gatekeeper.deduplicate_records(list(raw_records))
        lines: List[AvantMetreLine] = []
        rows: List[RebarScheduleRow] = []

        for rec in records:
            dims = self.gatekeeper.normalize_dimensions_to_meters(
                rec.label, rec.category, rec.raw_dimensions)
            qty = self.gatekeeper.effective_quantity(rec)
            sched, steel_kg = self.gatekeeper.parse_rebar_specs_to_schedule(
                rec, dims, qty)
            rows.extend(sched)

            if rec.category == "semelle":
                line_bp, line_ba = self.engine.compute_semelle_isolee(
                    label=rec.label, qty=qty,
                    A=dims.get("A", dims.get("a", 1.20)),
                    B=dims.get("B", dims.get("b", 1.20)),
                    h_patin=dims.get("h_patin", dims.get("H", dims.get("h", 0.35))),
                    h_glacis=dims.get("h_glacis", 0.0),
                    a_fut=dims.get("a_fut", 0.30),
                    b_fut=dims.get("b_fut", 0.30),
                )
                line_bp.grid_locations = list(rec.grid_axes)
                line_ba.grid_locations = list(rec.grid_axes)
                line_ba.acier_kg = steel_kg
                self.gatekeeper.cross_check(line_ba, rec)
                lines.extend([line_bp, line_ba])

            elif rec.category in ("poutre", "longrine", "chainage"):
                line = self.engine.compute_poutre_entre_nus(
                    label=rec.label, qty=qty,
                    portee_axe_axe=dims.get("L", dims.get("portee", 4.50)),
                    largeur_poteau_gauche=rec.support_left_m,
                    largeur_poteau_droite=rec.support_right_m,
                    b_poutre=dims.get("b", 0.25),
                    h_totale=dims.get("h", 0.45),
                    type_plancher=rec.type_plancher,
                    epaisseur_dalle=rec.epaisseur_dalle_m,
                    category=rec.category,
                )
                line.grid_locations = list(rec.grid_axes)
                line.acier_kg = steel_kg
                self.gatekeeper.cross_check(line, rec)
                lines.append(line)

            elif rec.category in ("poteau", "voile"):
                line = self.engine.compute_poteau(
                    label=rec.label, qty=qty,
                    b=dims.get("b", dims.get("e", 0.25)),
                    h=dims.get("h", dims.get("L", 0.30))
                    if rec.category == "voile" else dims.get("h", 0.30),
                    hauteur=dims.get("H", dims.get("hauteur", dims.get("L", 3.00)))
                    if rec.category == "voile"
                    else dims.get("L", dims.get("H", dims.get("hauteur", 3.00))),
                    category=rec.category,
                )
                line.grid_locations = list(rec.grid_axes)
                line.acier_kg = steel_kg
                self.gatekeeper.cross_check(line, rec)
                lines.append(line)

            elif rec.category == "dalle_pleine":
                L = dims.get("L", dims.get("A", 0.0))
                l_ = dims.get("B", dims.get("b", 0.0))
                e = dims.get("e", dims.get("ep", dims.get("h", 0.0)))
                line = AvantMetreLine(
                    prix_code=PRIX_CODES["dalle_pleine"],
                    category="dalle_pleine",
                    designation=f"Dalle pleine BA {rec.label} ép. {e:.2f} m",
                    unite="m3", nombre=qty, longueur=round(L, 3),
                    largeur=round(l_, 3), hauteur=round(e, 3),
                    quantite_partielle=round(qty * L * l_ * e, 3),
                    acier_kg=steel_kg, grid_locations=list(rec.grid_axes),
                    label=rec.label,
                )
                self.gatekeeper.cross_check(line, rec)
                lines.append(line)
            else:
                self.engine.anomalies.append(AuditAnomaly(
                    severity="AVERTISSEMENT", element_label=rec.label,
                    rule_code="CATEGORY_NOT_ROUTED",
                    message=f"Catégorie '{rec.category}' non métrée.",
                    suggested_fix="Reclasser l'élément (semelle, poteau, "
                                  "poutre, longrine, chainage, voile)."))

        self.records = records
        self.avant_metre_lines = lines
        self.rebar_schedule_rows = rows
        return lines, rows

    # ------------------------------------------------------------------
    def process_and_export(self, raw_records: List[ExtractedPlanRecord],
                           output_xlsx_path: Union[str, Path]) -> Path:
        lines, rows = self.compute(raw_records)
        crops = [(r.label, r.crop_image_bgr) for r in self.records
                 if r.crop_image_bgr is not None]

        wb = Workbook()
        wb.remove(wb.active)
        self.layout = {}
        self._build_sheet_avant_metre(wb, lines)
        self._build_sheet_nomenclature_acier(wb, rows)
        self._build_sheet_bordereau_dqe(wb)
        self._build_sheet_controle_bet(wb, crops)
        self._build_sheet_optimisation_12m(wb, rows)

        out = Path(output_xlsx_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(out))
        return out

    # ------------------------------------------------------------------
    # Feuille 1
    # ------------------------------------------------------------------
    def _build_sheet_avant_metre(self, wb: Workbook,
                                 lines: List[AvantMetreLine]) -> None:
        ws = wb.create_sheet(title=SHEET_NAMES[0])
        self._write_banner(
            ws, f"AVANT-MÉTRÉ DÉTAILLÉ BÉTON ARMÉ — {self.project_name}",
            "DGA / CPS Maroc / BAEL 91 mod. 99 — portées strictes entre nus",
            max_col=12)
        headers = ["N° Prix", "Famille", "Désignation des Ouvrages",
                   "Files d'Axes / Repérage", "Unité", "Nombre",
                   "Longueur (m)", "Largeur (m)", "Hauteur / Ép. (m)",
                   "Cube Partiel (m³)", "Témoin Python (m³)", "Écart"]
        self._write_table_headers(ws, 5, headers)

        start = 6
        for i, line in enumerate(lines):
            r = start + i
            axes = ", ".join(line.grid_locations) if line.grid_locations else "Vue en plan"
            vals = [line.prix_code, line.category, line.designation, axes,
                    line.unite, line.nombre, line.longueur, line.largeur,
                    line.hauteur, f"=ROUND(F{r}*G{r}*H{r}*I{r}, 3)",
                    line.quantite_partielle, f"=ROUND(J{r}-K{r}, 3)"]
            self._write_data_row(ws, r, vals, i % 2 == 1, {10, 12})
            for c in (7, 8, 9, 10, 11, 12):
                ws.cell(row=r, column=c).number_format = "0.000"

        end = max(start, start + len(lines) - 1)
        total_row = end + 1
        ws.cell(row=total_row, column=3, value="TOTAL GÉNÉRAL BÉTON (m³)").font = self.font_bold
        tot = ws.cell(row=total_row, column=10, value=f"=ROUND(SUM(J{start}:J{end}), 3)")
        tot.font = self.font_formula
        tot.number_format = "0.000"
        tk = ws.cell(row=total_row, column=11,
                     value=round(sum(l.quantite_partielle for l in lines), 3))
        tk.number_format = "0.000"
        for c in range(1, 13):
            ws.cell(row=total_row, column=c).fill = self.fill_total
            ws.cell(row=total_row, column=c).border = self.border_total
        ws.freeze_panes = "A6"
        self._autofit(ws)
        self.layout["avant_metre"] = {"start": start, "end": end, "total_row": total_row}

    # ------------------------------------------------------------------
    # Feuille 2
    # ------------------------------------------------------------------
    def _build_sheet_nomenclature_acier(self, wb: Workbook,
                                        rows: List[RebarScheduleRow]) -> None:
        ws = wb.create_sheet(title=SHEET_NAMES[1])
        self._write_banner(
            ws, f"NOMENCLATURE DES ACIERS HA FeE500 — {self.project_name}",
            "RPS 2000 v2011 : crochets 135°, ancrages 15Φ, recouvrements 50Φ par barre de 12 m",
            max_col=8 + len(self.DIAMETERS_COLS))
        headers = ["Ouvrage", "Catégorie", "Repère Barre", "Φ (mm)",
                   "Nb Ouvrages", "Nb Barres / Ouvr.", "Long. Développée (m)",
                   "Long. Totale (ml)"] + [f"HA{d} (kg)" for d in self.DIAMETERS_COLS]
        self._write_table_headers(ws, 5, headers)

        start = 6
        for i, it in enumerate(rows):
            r = start + i
            vals: List[Any] = [it.element_label, it.category, it.bar_mark,
                               it.diameter_mm, it.nb_elements, it.nb_bars_per_elem,
                               it.developed_length_m, f"=ROUND(E{r}*F{r}*G{r}, 2)"]
            for d in self.DIAMETERS_COLS:
                vals.append(f"=IF($D{r}={d}, ROUND($H{r}*{MOROCCAN_REBAR_CATALOG[d]}, 2), 0)")
            self._write_data_row(ws, r, vals, i % 2 == 1, set(range(8, 9 + len(self.DIAMETERS_COLS))))
            for c in range(7, 9 + len(self.DIAMETERS_COLS)):
                ws.cell(row=r, column=c).number_format = "0.00"

        end = max(start, start + len(rows) - 1)
        tot_row = end + 1
        ws.cell(row=tot_row, column=3, value="SOUS-TOTAL POIDS PAR DIAMÈTRE (kg)").font = self.font_bold
        first_dia_col = 9
        last_dia_col = 8 + len(self.DIAMETERS_COLS)
        for i in range(len(self.DIAMETERS_COLS)):
            col = first_dia_col + i
            L = get_column_letter(col)
            c = ws.cell(row=tot_row, column=col, value=f"=ROUND(SUM({L}{start}:{L}{end}), 2)")
            c.font = self.font_formula
            c.number_format = "0.00"
        for c in range(1, last_dia_col + 1):
            ws.cell(row=tot_row, column=c).fill = self.fill_total
            ws.cell(row=tot_row, column=c).border = self.border_total

        grand_row = tot_row + 1
        ws.cell(row=grand_row, column=3, value="POIDS TOTAL ACIER HA FeE500 (kg)").font = self.font_bold
        g = ws.cell(row=grand_row, column=first_dia_col,
                    value=f"=ROUND(SUM({get_column_letter(first_dia_col)}{tot_row}:"
                          f"{get_column_letter(last_dia_col)}{tot_row}), 2)")
        g.font = self.font_formula
        g.fill = self.fill_ok
        g.number_format = "0.00"
        ws.freeze_panes = "A6"
        self._autofit(ws)
        self.layout["acier"] = {
            "start": start, "end": end, "tot_row": tot_row,
            "grand_row": grand_row, "first_dia_col": first_dia_col,
            "last_dia_col": last_dia_col,
            "grand_cell": f"{get_column_letter(first_dia_col)}{grand_row}",
        }

    # ------------------------------------------------------------------
    # Feuille 3
    # ------------------------------------------------------------------
    def _build_sheet_bordereau_dqe(self, wb: Workbook) -> None:
        ws = wb.create_sheet(title=SHEET_NAMES[2])
        self._write_banner(
            ws, f"BORDEREAU DES PRIX — DQE — {self.project_name}",
            "Quantités liées dynamiquement aux feuilles 1 et 2 (MAD HT / TVA / TTC)",
            max_col=7)
        headers = ["N° Prix", "Désignation des Prestations (CPS Maroc)",
                   "Unité", "Quantité", "Prix Unitaire HT (MAD)",
                   "Montant Total HT (MAD)", "Référence Normative"]
        self._write_table_headers(ws, 5, headers)

        s1, s2 = SHEET_NAMES[0], SHEET_NAMES[1]
        acier = self.layout["acier"]
        start = 6
        for i, (code, desc, unit, pu, ref) in enumerate(BPU_ITEMS):
            r = start + i
            if code.startswith("B."):
                qty = (f"=ROUND(SUMIF('{s1}'!$A$6:$A${DATA_RANGE_END}, A{r}, "
                       f"'{s1}'!$J$6:$J${DATA_RANGE_END}), 3)")
            else:
                qty = f"=ROUND('{s2}'!{acier['grand_cell']}, 2)"
            vals = [code, desc, unit, qty, pu, f"=ROUND(D{r}*E{r}, 2)", ref]
            self._write_data_row(ws, r, vals, i % 2 == 1, {4, 6})
            ws.cell(row=r, column=4).number_format = "0.000"
            ws.cell(row=r, column=5).number_format = "#,##0.00"
            ws.cell(row=r, column=6).number_format = "#,##0.00"

        end = start + len(BPU_ITEMS) - 1
        t_ht, t_tva, t_ttc = end + 1, end + 2, end + 3
        ws.cell(row=t_ht, column=2, value="TOTAL GÉNÉRAL TRAVAUX HT (MAD)").font = self.font_bold
        ws.cell(row=t_ht, column=6, value=f"=ROUND(SUM(F{start}:F{end}), 2)").font = self.font_formula
        ws.cell(row=t_tva, column=2, value=f"TVA ({self.tva_rate*100:.0f} %)").font = self.font_bold
        ws.cell(row=t_tva, column=6, value=f"=ROUND(F{t_ht}*{self.tva_rate}, 2)").font = self.font_formula
        ws.cell(row=t_ttc, column=2, value="TOTAL GÉNÉRAL TRAVAUX TTC (MAD)").font = self.font_bold
        ws.cell(row=t_ttc, column=6, value=f"=ROUND(F{t_ht}+F{t_tva}, 2)").font = self.font_formula
        for rr in (t_ht, t_tva, t_ttc):
            ws.cell(row=rr, column=6).number_format = "#,##0.00"
            for c in range(1, 8):
                ws.cell(row=rr, column=c).fill = self.fill_total
                ws.cell(row=rr, column=c).border = self.border_total
        self._autofit(ws)
        self.layout["dqe"] = {"start": start, "end": end, "t_ht": t_ht,
                              "t_tva": t_tva, "t_ttc": t_ttc}

    # ------------------------------------------------------------------
    # Feuille 4
    # ------------------------------------------------------------------
    def _build_sheet_controle_bet(self, wb: Workbook,
                                  crops: List[Tuple[str, Any]]) -> None:
        ws = wb.create_sheet(title=SHEET_NAMES[3])
        self._write_banner(
            ws, f"CONTRÔLE QUALITÉ BET MAROC — RATIOS RPS 2011 — {self.project_name}",
            "Vérification automatique par famille d'ouvrages et journal des anomalies",
            max_col=8)
        headers = ["Catégorie d'Ouvrage", "Volume Béton Cumulé (m³)",
                   "Poids Acier Cumulé (kg)", "Ratio Réel (kg/m³)",
                   "Ratio Min BET Maroc", "Ratio Max BET Maroc",
                   "Statut Normatif", "Observation / Action Requise"]
        self._write_table_headers(ws, 5, headers)

        s1, s2 = SHEET_NAMES[0], SHEET_NAMES[1]
        acier = self.layout["acier"]
        start = 6
        for i, cat in enumerate(self.CONTROL_CATEGORIES):
            r = start + i
            min_r, max_r = MOROCCAN_STEEL_RATIO_BOUNDS[cat]
            vol = (f"=ROUND(SUMIF('{s1}'!$B$6:$B${DATA_RANGE_END}, A{r}, "
                   f"'{s1}'!$J$6:$J${DATA_RANGE_END}), 3)")
            terms = []
            for k in range(len(self.DIAMETERS_COLS)):
                L = get_column_letter(acier["first_dia_col"] + k)
                terms.append(f"SUMIF('{s2}'!$B$6:$B${DATA_RANGE_END}, A{r}, "
                             f"'{s2}'!${L}$6:${L}${DATA_RANGE_END})")
            steel = f"=ROUND({'+'.join(terms)}, 2)"
            ratio = f"=IF(B{r}=0, 0, ROUND(C{r}/B{r}, 1))"
            status = (f'=IF(B{r}=0, "NON APPLICABLE", '
                      f'IF(AND(D{r}>=E{r}, D{r}<=F{r}), "CONFORME BET MAROC", '
                      f'"ALERTE HORS PLAGE RPS 2011"))')
            obs = (f'=IF(B{r}=0, "-", IF(D{r}<E{r}, '
                   f'"Sous-ferraillage ou oubli de barres à vérifier", '
                   f'IF(D{r}>F{r}, "Sur-ferraillage ou confusion d\'unité (cm/m) à vérifier", '
                   f'"Ratio validé")))')
            vals = [cat, vol, steel, ratio, min_r, max_r, status, obs]
            self._write_data_row(ws, r, vals, i % 2 == 1, {2, 3, 4, 7, 8})
            ws.cell(row=r, column=2).number_format = "0.000"
            ws.cell(row=r, column=3).number_format = "0.00"
            ws.cell(row=r, column=4).number_format = "0.0"
        end = start + len(self.CONTROL_CATEGORIES) - 1
        rng = f"G{start}:G{end}"
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'$G{start}="CONFORME BET MAROC"'], fill=self.fill_ok))
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'$G{start}="ALERTE HORS PLAGE RPS 2011"'], fill=self.fill_err))

        # Journal des anomalies
        a_title = end + 3
        ws.cell(row=a_title, column=1,
                value="JOURNAL DES ANOMALIES — CONTRÔLEUR TOPOLOGIQUE & OCR").font = self.font_bold
        self._write_table_headers(ws, a_title + 1, [
            "Sévérité", "Ouvrage Concerné", "Code Règle",
            "Description Technique de l'Anomalie", "Correction Préconisée"])
        a_start = a_title + 2
        anomalies = self.engine.anomalies
        if not anomalies:
            c = ws.cell(row=a_start, column=1,
                        value="AUCUNE ANOMALIE — Concordance Coffrage / "
                              "Nomenclature et ratios RPS 2011 validés.")
            c.font = self.font_bold
            c.fill = self.fill_ok
        for i, an in enumerate(anomalies):
            r = a_start + i
            self._write_data_row(ws, r, [an.severity, an.element_label,
                                         an.rule_code, an.message,
                                         an.suggested_fix or "Vérification manuelle requise"])
            ws.cell(row=r, column=1).fill = (
                self.fill_err if an.severity == "BLOQUANT" else self.fill_warn)
            ws.cell(row=r, column=1).font = self.font_bold

        # Vignettes OpenCV (preuves d'extraction)
        img_row = a_start + max(1, len(anomalies)) + 2
        if crops:
            ws.cell(row=img_row, column=1,
                    value="EXTRAITS VISUELS DU PLAN (PREUVES D'AUDIT COMPUTER VISION)"
                    ).font = self.font_bold
            cur = img_row + 1
            for label, img_bgr in crops[:12]:
                ws.cell(row=cur, column=1, value=f"Repère : {label}").font = self.font_bold
                xl_img = self._to_xl_image(img_bgr, max_h=110)
                if xl_img is not None:
                    ws.row_dimensions[cur].height = 90
                    ws.add_image(xl_img, f"B{cur}")
                cur += 1
        self._autofit(ws)
        self.layout["controle"] = {"start": start, "end": end,
                                   "anom_start": a_start, "n_anomalies": len(anomalies)}

    @staticmethod
    def _to_xl_image(img_bgr, max_h: int = 110):
        try:
            import cv2
            from openpyxl.drawing.image import Image as XLImage
            ok, buf = cv2.imencode(".png", img_bgr)
            if not ok:
                return None
            xl = XLImage(io.BytesIO(buf.tobytes()))
            scale = max_h / max(xl.height, 1)
            xl.height = int(xl.height * scale)
            xl.width = int(xl.width * scale)
            return xl
        except Exception:
            return None

    # ------------------------------------------------------------------
    # Feuille 5
    # ------------------------------------------------------------------
    def _build_sheet_optimisation_12m(self, wb: Workbook,
                                      rows: List[RebarScheduleRow]) -> None:
        ws = wb.create_sheet(title=SHEET_NAMES[4])
        self._write_banner(
            ws, f"OPTIMISATION DE COUPE — BARRES COMMERCIALES 12.00 m — {self.project_name}",
            "Algorithme First-Fit Decreasing (FFD) : minimisation des chutes chantier",
            max_col=7)
        headers = ["Diamètre Nominal", "Nb Barres de 12.00 m à Commander",
                   "Longueur Utile Calculée (ml)", "Longueur Totale Commandée (ml)",
                   "Chutes Totales (ml)", "Taux de Chute (%)",
                   "Poids Total à Commander (kg)"]
        self._write_table_headers(ws, 5, headers)
        summary = optimize_12m_stock_bars(rows, 12.00)
        start = 6
        for i, (dia, st) in enumerate(summary.items()):
            r = start + i
            vals = [f"HA {dia}", st["nb_stock_bars_12m"], st["useful_length_m"],
                    f"=B{r}*12", f"=ROUND(D{r}-C{r}, 2)",
                    f"=IF(D{r}=0, 0, ROUND((E{r}/D{r})*100, 2))",
                    f"=ROUND(D{r}*{MOROCCAN_REBAR_CATALOG.get(dia, 0.0)}, 2)"]
            self._write_data_row(ws, r, vals, i % 2 == 1, {4, 5, 6, 7})
            for c in (3, 4, 5, 6, 7):
                ws.cell(row=r, column=c).number_format = "0.00"
        end = max(start, start + len(summary) - 1)
        tot = end + 1
        ws.cell(row=tot, column=1, value="TOTAL COMMANDE").font = self.font_bold
        for c, L in ((2, "B"), (4, "D"), (5, "E"), (7, "G")):
            cell = ws.cell(row=tot, column=c, value=f"=ROUND(SUM({L}{start}:{L}{end}), 2)")
            cell.font = self.font_formula
            cell.number_format = "0.00"
        for c in range(1, 8):
            ws.cell(row=tot, column=c).fill = self.fill_total
            ws.cell(row=tot, column=c).border = self.border_total
        self._autofit(ws)
        self.layout["optim"] = {"start": start, "end": end, "tot_row": tot,
                                "summary": summary}

    # ------------------------------------------------------------------
    # Utilitaires de style
    # ------------------------------------------------------------------
    def _write_banner(self, ws, title: str, subtitle: str, max_col: int) -> None:
        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=max_col)
        c1 = ws.cell(row=1, column=1, value=title)
        c1.font = self.font_title
        c1.fill = self.fill_navy
        c1.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max_col)
        c2 = ws.cell(row=3, column=1, value=f"{subtitle} | Généré par {self.bet_name}")
        c2.font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
        c2.fill = self.fill_subhdr
        c2.alignment = Alignment(horizontal="center", vertical="center")

    def _write_table_headers(self, ws, row: int, headers: List[str]) -> None:
        ws.row_dimensions[row].height = 30
        for col, text in enumerate(headers, start=1):
            c = ws.cell(row=row, column=col, value=text)
            c.font = self.font_header
            c.fill = self.fill_navy
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = self.border_cell

    def _write_data_row(self, ws, row: int, values: List[Any],
                        zebra: bool = False, formula_cols: Optional[set] = None) -> None:
        formula_cols = formula_cols or set()
        ws.row_dimensions[row].height = 19
        for col, val in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=val)
            c.font = self.font_formula if col in formula_cols else self.font_normal
            if zebra:
                c.fill = self.fill_zebra
            c.border = self.border_cell
            right = isinstance(val, (int, float)) or col in formula_cols
            c.alignment = Alignment(horizontal="right" if right else "left",
                                    vertical="center")

    @staticmethod
    def _autofit(ws) -> None:
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = 12
            for cell in col:
                if cell.row <= 3:
                    continue
                s = str(cell.value or "")
                width = max(width, 14 if s.startswith("=") else min(len(s) + 4, 48))
            ws.column_dimensions[letter].width = width
